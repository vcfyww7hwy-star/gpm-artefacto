# -*- coding: utf-8 -*-
"""r2_report.py REF_v20_calc.xlsx NEW_v30_calc.xlsx SALIDA.txt — informe de deltas R2 (v3.0 política B frente a la v2.0) mapeado a los M del doc 13.
 1. KPI de los cuatro casos antes/después (nombres X_/C_/B_/F_).
 2. Puente BR0 → BR5 del Motor v3.0 (atribución del cambio del Base a cada parámetro nuevo: M-a escalación, M-d disponibilidad, M-c reemplazo,
    M-e meses de construcción/IDC, M-f tasa del accionista); los M en neutro (M-b pool, M-g payback robusto, M-h peaje kW, degradación, desmantelamiento) no mueven nada.
 3. Nombres definidos que cambian, agrupados por familia → M responsable(s); Motor: celdas distintas por bloque → M.
Complementa a regress30 (que en modo v30 sólo lista los nombres distintos)."""
import sys, re
from collections import OrderedDict, Counter
from openpyxl import load_workbook

ref_path, new_path, out_path = sys.argv[1], sys.argv[2], sys.argv[3]
R = load_workbook(ref_path, data_only=True); RF = load_workbook(ref_path, data_only=False)
N = load_workbook(new_path, data_only=True); NF = load_workbook(new_path, data_only=False)
MOTOR = "Motor_Sens"


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def cells(wbv, wbf, nm):
    dn = wbf.defined_names.get(nm)
    if dn is None:
        return None
    ref = dn.attr_text.replace("'", "").replace("$", "")
    sh, rng = ref.split("!")
    ws = wbv[sh]
    if ":" in rng:
        return [c.value for row in ws[rng] for c in (row if isinstance(row, tuple) else (row,))]
    return [ws[rng].value]


def one(wbv, wbf, nm):
    v = cells(wbv, wbf, nm)
    return v[0] if v else None


# familias de nombres → M responsable (orden: la primera regla que casa)
FAM = [
    (r"^(P50|P90)_(?!Ahorro1)|^Ahorro_kWh$|_Row$", "resultado combinado (M-a + M-d + M-c + M-e)"),
    (r"_E1$|_Ahorro1$|^Red_Factura$", "M-d disponibilidad (energía × 0,98) → ahorro"),
    (r"(Carga_Exergy|Nominal_Exergy|VAN_Gerencia|VAN_Exergy|TIR_Exergy|_VANX$)", "Exergy: M-a (fee de gerencia ∝ CAPEX) + M-d (carga ∝ ahorro)"),
    (r"^(E_|Energia|Yield|Prod|Eval|E_Val|Cobertura|NoRec|Ahorro|Reduccion|P50_|P90_|Frac_Peaje)", "M-d disponibilidad (energía × 0,98) → ahorro"),
    (r"(CAPEX|IVA|Factor_Caso|Factor_Escalacion|Anios_Precios|Subtotal|Arancel|ISD|Fee_Gerencia_USD|Contingencia_USD|Nacionaliz|Depreciable|Dep_|DedAd|Escala_)", "M-a escalación del CAPEX (× fEsc)"),
    (r"(Reemplazo|Krep|Rep_|Residual|Decom|Desmantel)", "M-c reemplazo de inversores / desmantelamiento"),
    (r"(IDC|Meses_Construccion|Deuda|Cuota|Interes|Amort|Servicio|DSCR|CFADS|Aporte)", "M-e construcción 21 meses (IDC) + M-a (deuda ∝ CAPEX)"),
    (r"(VANeq|VAN_eq|VAN_Equity|Tasa_Descuento_Equity)", "M-f VAN del accionista a su tasa (12 %) + todo lo anterior"),
    (r"(TIR|VAN|PB|Payback|LCOE|Estado|N_|Tornado|Rank|Max|Min|Piso|Puente|Matriz|Lev|Sweep|Deuda_Max)", "resultado combinado (M-a + M-d + M-c + M-e)"),
]


def fam(nm):
    for pat, m in FAM:
        if re.search(pat, nm):
            return m
    return "otros (revisar)"


out = []
P = out.append
P("R2 · Informe de deltas v3.0 (política B) frente a v2.0 — mapeado a los M del doc 13")
P(f"REF = {ref_path}\nNEW = {new_path}\n")
# 1. KPI
P("1 · KPI de los cuatro casos (v2.0 → v3.0)")
P(f"  {'caso':12s} {'TIR':>17s} {'VAN [USD]':>26s} {'TIR acc.':>17s} {'VAN acc. [USD]':>26s} {'DSCR mín':>13s} {'PB':>11s} {'LCOE':>13s}")
for tag, nm in (("X", "Custom"), ("C", "Conservador"), ("B", "Base"), ("F", "Favorable")):
    row = [nm]
    for key, f in (("TIR", "{:.2%}"), ("VAN", "{:,.0f}"), ("TIReq", "{:.2%}"), ("VANeq", "{:,.0f}"), ("DSCR", "{:.2f}"), ("PB", "{:.1f}"), ("LCOE", "{:.1f}")):
        a, b = one(R, RF, f"{tag}_{key}"), one(N, NF, f"{tag}_{key}")
        fa = f.format(a) if is_num(a) else str(a); fb = f.format(b) if is_num(b) else str(b)
        row.append(f"{fa} → {fb}")
    P(f"  {row[0]:12s} {row[1]:>17s} {row[2]:>26s} {row[3]:>17s} {row[4]:>26s} {row[5]:>13s} {row[6]:>11s} {row[7]:>13s}")
# 2. puente (Motor NEW, casos BR0–BR5 por nombre de columna)
wn = N[MOTOR]
case_col = {wn.cell(row=4, column=c).value: c for c in range(2, wn.max_column + 1) if wn.cell(row=4, column=c).value}
rows = {}
for r in range(5, wn.max_row + 1):
    v = wn.cell(row=r, column=1).value
    if isinstance(v, str):
        rows[v] = r
LAB = {"TIR": "TIR del proyecto", "VAN": "VAN del proyecto", "TIR_eq": "TIR del accionista", "VAN_eq": "VAN del accionista", "DSCR_min": "DSCR mínimo", "PB": "Payback", "LCOE": "LCOE"}
brs = [k for k in case_col if k.startswith("BR")]
brs.sort(key=lambda k: int(re.search(r"\d+", k).group()))
P("\n2 · Puente BR0 → BR5 (Motor v3.0; cada escalón enciende un parámetro con los anteriores activos; BR5 ≡ Base)")
STEP_M = {"BR0": "definición v2.0 (todo en neutro)", "BR1": "M-a escalación del CAPEX", "BR2": "M-d disponibilidad", "BR3": "M-c reemplazo de inversores", "BR4": "M-e construcción (IDC × meses/12)", "BR5": "M-f tasa del accionista (sólo VAN acc.)"}
out_labels = {}
for lab, r in rows.items():
    for key in LAB:
        pass
# localizar filas de salida por etiqueta parcial
def find_row(keys):
    for lab, r in rows.items():
        if any(lab.lower().startswith(k.lower()) for k in keys):
            return r
    return None
ROWK = {"TIR": ["TIR proyecto", "TIR del proyecto"], "VAN": ["VAN proyecto", "VAN del proyecto"], "TIR_eq": ["TIR accionista", "TIR del accionista", "TIR equity"], "VAN_eq": ["VAN accionista", "VAN del accionista", "VAN equity"],
        "DSCR_min": ["DSCR mín", "DSCR min"], "PB": ["Payback"], "LCOE": ["LCOE"]}
rr = {k: find_row(v) for k, v in ROWK.items()}
P(f"  {'escalón':6s} {'M':44s} " + " ".join(f"{LAB[k]:>18s}" for k in ROWK))
prev = None
for br in brs:
    c = case_col[br]
    vals = {k: (wn.cell(row=rr[k], column=c).value if rr[k] else None) for k in ROWK}
    txt = []
    for k in ROWK:
        v = vals[k]
        if not is_num(v):
            txt.append(f"{str(v):>18s}"); continue
        d = "" if prev is None or not is_num(prev.get(k)) else (f" ({(v - prev[k]) * 100:+.2f} pp)" if k in ("TIR", "TIR_eq") else "")
        s = f"{v:.2%}" if k in ("TIR", "TIR_eq") else (f"{v:,.0f}" if k in ("VAN", "VAN_eq") else f"{v:.2f}")
        txt.append(f"{s + d:>18s}")
    P(f"  {br:6s} {STEP_M.get(br, ''):44s} " + " ".join(txt))
    prev = vals
# 3. nombres
P("\n3 · Nombres definidos que cambian (REF ∩ NEW), agrupados por M")
groups = OrderedDict()
n_same = 0
for nm in sorted(set(RF.defined_names) & set(NF.defined_names)):
    a, b = cells(R, RF, nm), cells(N, NF, nm)
    if a is None or b is None or len(a) != len(b):
        continue
    dmax = 0.0; rel = 0.0
    for va, vb in zip(a, b):
        if is_num(va) and is_num(vb):
            d = abs(va - vb); dmax = max(dmax, d)
            if abs(va) > 1e-9: rel = max(rel, d / abs(va))
        elif is_num(va) != is_num(vb):
            dmax = float("inf")
    if dmax == 0:
        n_same += 1; continue
    groups.setdefault(fam(nm), []).append((nm, dmax, rel, a[0] if len(a) == 1 else None, b[0] if len(b) == 1 else None))
P(f"  nombres iguales: {n_same} · distintos: {sum(len(v) for v in groups.values())}")
for m, items in groups.items():
    P(f"  — {m}: {len(items)}")
    for nm, dmax, rel, a0, b0 in sorted(items):
        if a0 is not None and is_num(a0) and is_num(b0):
            P(f"      {nm:34s} {a0:>16,.6g} → {b0:>16,.6g}   (Δ {b0 - a0:+,.6g}{'' if abs(a0) < 1e-9 else f' · {(b0 - a0) / abs(a0):+.1%}'})")
        else:
            P(f"      {nm:34s} rango: Δmax {dmax:,.6g}" + (f" · rel {rel:.1%}" if rel else ""))
# 4. Motor por bloque
P("\n4 · Motor: celdas anuales distintas por bloque (casos comunes, por etiqueta) → M")
wr = R[MOTOR]
ref_cases = {wr.cell(row=4, column=c).value: c for c in range(2, wr.max_column + 1) if wr.cell(row=4, column=c).value}
def motor_rows(ws):
    rows_ = {}; block = None
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v is None: continue
        if isinstance(v, str) and v.startswith("Bloque "):
            block = v.split(" ")[1]; continue
        if block is not None and is_num(v):
            rows_[(block, int(v))] = r
        elif isinstance(v, str) and r >= 5:
            rows_[v] = r
    return rows_
rrf, rnw = motor_rows(wr), motor_rows(wn)
BLOCK_M = {"E": "M-d", "Eval": "M-d", "Ahorro": "M-d", "Peaje": "M-d", "OPEX": "—", "EBITDA": "M-d", "Dep": "M-a/M-c", "Part_u": "M-a/M-d/M-c", "IR_u": "M-a/M-d/M-c", "Terr": "—", "FCF_u": "todos", "Cum_u": "todos",
           "Int": "M-a/M-e", "Amort": "M-a/M-e", "Part_l": "todos", "IR_l": "todos", "CFADS": "todos", "EQ": "todos", "DSCR": "todos", "DF": "M-a/M-e", "Ux": "M-c (Exergy)", "Ix": "M-c", "Tx": "M-c", "Fx": "M-c", "Gx": "todos"}
cnt = Counter(); tot = Counter(); scal = Counter()
for nm, c in ref_cases.items():
    if nm not in case_col: continue
    c2 = case_col[nm]
    for k, r in rrf.items():
        if k not in rnw: continue
        va = wr.cell(row=r, column=c).value; vb = wn.cell(row=rnw[k], column=c2).value
        if not is_num(va) or not is_num(vb): continue
        if isinstance(k, tuple):
            tot[k[0]] += 1
            if abs(va - vb) > 0: cnt[k[0]] += 1
        else:
            if abs(va - vb) > 0: scal[k] += 1
for blk in tot:
    P(f"  {blk:8s} {cnt[blk]:6d} / {tot[blk]:6d} distintas   → {BLOCK_M.get(blk, '?')}")
P("  escalares/salidas distintas por etiqueta (n.º de casos): " + "; ".join(f"{k} {v}" for k, v in sorted(scal.items(), key=lambda x: -x[1])[:40]))
open(out_path, "w", encoding="utf-8").write("\n".join(out) + "\n")
print("\n".join(out[:40]))
print(f"… escrito en {out_path}")

# -*- coding: utf-8 -*-
"""Modelo sombra v2.0 (Python puro, código independiente del generador): energía × recorte, CAPEX por drivers o FIJO $/Wp (kfix),
terreno según comprador, IDC de ambos tramos, deuda con guardas, fiscal SALELGI, Exergy y grupo. Reproduce las 99 columnas de Motor_Sens
leyendo los parámetros de cada caso (filas 5–24) y los inputs compartidos del libro. El factor de caso NO se lee de 05 (v1.2 leía Factor_Caso):
CAPEX = fK × bottom-up, o kfix × kWp × 1000 si kfix > 0, exactamente como el Motor v2.0.
Uso: python3 shadow20.py libro_valores.xlsx salida.csv"""
import sys, csv, math
from datetime import date
import numpy as np
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as gl
from build_core import TS, C, PARAM_ROWS, SCAL_ROWS, OUT_ROWS, CAPEX as KL, ENERGIA as EL, FLUJO, EXERGY


def irr(cfs):
    cfs = np.array(cfs, dtype=float)
    f = lambda r: np.sum(cfs / (1 + r) ** np.arange(len(cfs)))
    lo, hi = -0.99, 1.0
    if f(lo) * f(hi) > 0:
        return float("nan")
    for _ in range(300):
        mid = (lo + hi) / 2
        if f(lo) * f(mid) <= 0: hi = mid
        else: lo = mid
    return (lo + hi) / 2


def npv0(rate, cfs):  # cfs[0]=t−1 capitalizado; cfs[1]=t0; resto descontado
    return cfs[0] * (1 + rate) + cfs[1] + sum(cf / (1 + rate) ** t for t, cf in zip(range(1, len(cfs) - 1), cfs[2:]))


def edate(d, months):
    y = d.year + (d.month - 1 + months) // 12; m = (d.month - 1 + months) % 12 + 1
    return date(y, m, min(d.day, 28))


class Book:
    def __init__(self, path):
        self.wb = load_workbook(path, data_only=True)
        self.names = {n: d.attr_text for n, d in self.wb.defined_names.items()}

    def v(self, nm):
        ref = self.names[nm].split("!"); sh = ref[0].strip("'"); cell = ref[1].replace("$", "")
        if ":" in cell:
            ws = self.wb[sh]; a, b = cell.split(":")
            return [c.value for row in ws[a:b] for c in row]
        return self.wb[sh][cell].value


def load_inputs(B):
    P = {}
    for nm in ["Horizonte", "Tasa_Descuento", "Fase_m1", "Tasa_IVA", "FODINFA_Pct", "ISD_Pct", "Contingencia_Pct", "Contingencia_Frac_IVA", "Fee_Gerencia_Pct", "Asignacion_Compartida",
               "Fee_OM_kWp", "Seguro_kWp", "Renta_Terreno_ha", "Tributos_Locales", "Escalacion_OPEX", "Precio_Terreno_ha", "Costos_Transaccion_Terreno_Pct", "Predial_Terreno", "Residual_Terreno_Pct", "Apreciacion_Terreno",
               "Tasa_IR", "Tasa_Participacion", "Vida_Fiscal_Equipos", "Vida_Fiscal_Civil", "Pct_Elegible_DedAd", "Ingresos_SALELGI", "Tope_DedAd_Pct", "IDC_Frac_Tramo0",
               "Costo_Gerencia_Pct", "Costo_OM_Exergy_kWp", "Tasa_Efectiva_Exergy", "Potencia_Ref", "Ratio_Ref", "Densidad_MWp_ha", "Exponente_Escala", "Crecimiento_Consumo",
               "Tarifa_Evitable", "Consumo_Anual", "Pct_CAPEX_Civil", "Tasa_Efectiva"]:
        P[nm] = B.v(nm)
    P["escudo"] = B.v("Escudo_Negativo") == "Sí"
    P["cod"] = B.v("Fecha_COD"); P["fpeaje"] = B.v("Fecha_Peaje")
    P["cod"] = P["cod"].date() if hasattr(P["cod"], "date") else P["cod"]; P["fpeaje"] = P["fpeaje"].date() if hasattr(P["fpeaje"], "date") else P["fpeaje"]
    P["Y50"] = [x for x in B.v("Y_P50")]; P["Y90"] = [x for x in B.v("Y_P90")]
    P["cr_ratio"] = B.v("CR_Ratio"); P["cr_loss"] = B.v("CR_Loss")
    # rubros de 05_CAPEX: costo base D, %comp E, %ext G, arancel I, IVA O, drivers S T U
    ws = B.wb["05_CAPEX"]
    P["rubros"] = [dict(cost=ws[f"D{r}"].value, comp=ws[f"E{r}"].value, ext=ws[f"G{r}"].value, ar=ws[f"I{r}"].value, iva=ws[f"O{r}"].value, wp=ws[f"S{r}"].value, wac=ws[f"T{r}"].value, wf=ws[f"U{r}"].value)
                   for r in range(KL["r0"], KL["r0"] + 9)]
    P["iva_fee"] = ws[f"O{KL['fee']}"].value
    return P


def loss(P, r):
    xs, ys = P["cr_ratio"], P["cr_loss"]
    r = min(max(r, xs[0]), xs[-1])
    for i in range(len(xs) - 1):
        if xs[i] <= r <= xs[i + 1]:
            return ys[i] + (r - xs[i]) * (ys[i + 1] - ys[i]) / (xs[i + 1] - xs[i])
    return ys[-1]


def capex(P, fK, cont, Pkw, ratio, kfix=0.0):
    """CAPEX industrial sin IVA e IVA para un caso: costo real bottom-up × fK, o kfix $/Wp × kWp × 1000 si kfix > 0 (el IVA y el
    subtotal EPC escalan en la misma proporción), con escala por drivers."""
    a = (Pkw / P["Potencia_Ref"]) ** (1 - P["Exponente_Escala"]); b = ((Pkw / ratio) / (P["Potencia_Ref"] / P["Ratio_Ref"])) ** (1 - P["Exponente_Escala"])
    cap = iva = 0.0
    for rb in P["rubros"]:
        S = rb["cost"] * (1 - rb["comp"] * (1 - P["Asignacion_Compartida"]))
        esc = rb["wp"] * a + rb["wac"] * b + rb["wf"]
        F = S * esc  # costo a factor 1 escalado
        vext = F * rb["ext"]; ar = vext * rb["ar"]; fod = vext * P["FODINFA_Pct"]; isd = vext * P["ISD_Pct"]
        cap += F + fod + (0 if cont else ar + isd)
        iva += (F + fod + (0 if cont else ar)) * rb["iva"]
    contg = P["Contingencia_Pct"] * cap
    iva += contg * P["Tasa_IVA"] * P["Contingencia_Frac_IVA"]
    sub = cap + contg
    fee = P["Fee_Gerencia_Pct"] * sub
    iva += fee * P["iva_fee"]
    base = sub + fee
    f = (kfix * Pkw * 1000 / base) if (kfix and kfix > 0) else fK
    return base * f, iva * f, sub * f


def run_case(P, pc):
    H = int(P["Horizonte"]); T = list(range(-1, H + 1)); r = P["Tasa_Descuento"]
    Pkw, ratio, terr, finT, fPre = pc["P"], pc["ratio"], int(pc["terr"]), int(pc["finT"]), pc["fPre"]
    frec = (1 - loss(P, ratio)) / (1 - loss(P, P["Ratio_Ref"]))
    ha = Pkw / 1000 / P["Densidad_MWp_ha"]
    K, IVA, sub = capex(P, pc["fK"], int(pc["cont"]) == 1, Pkw, ratio, kfix=(pc.get("kfix") or 0.0))
    recup = int(pc["iva"]) == 1; part_on = int(pc["part"]) == 1; esc = P["escudo"]
    Kdep = K + (0 if recup else IVA)
    dep_eq = Kdep * (1 - P["Pct_CAPEX_Civil"]) / P["Vida_Fiscal_Equipos"]; dep_civ = Kdep * P["Pct_CAPEX_Civil"] / P["Vida_Fiscal_Civil"]
    dedad = min(Kdep * P["Pct_Elegible_DedAd"] / P["Vida_Fiscal_Equipos"], P["Tope_DedAd_Pct"] * P["Ingresos_SALELGI"])
    terr_cost = P["Precio_Terreno_ha"] * fPre * ha * (1 + P["Costos_Transaccion_Terreno_Pct"])
    resid_g = P["Precio_Terreno_ha"] * fPre * ha * P["Residual_Terreno_Pct"] * (1 + P["Apreciacion_Terreno"]) ** (H + 1)
    te_owner = P["Tasa_Efectiva"] if terr else P["Tasa_Efectiva_Exergy"]
    resid_net = resid_g - max(0, resid_g - terr_cost) * te_owner
    opex1 = (P["Fee_OM_kWp"] + P["Seguro_kWp"]) * Pkw + (P["Predial_Terreno"] if terr else P["Renta_Terreno_ha"] * ha) + P["Tributos_Locales"]
    fase = P["Fase_m1"]
    Y = P["Y50"] if int(pc["scen"]) == 1 else P["Y90"]
    def frac_peaje(t):
        a_ = edate(P["cod"], 12 * (t - 1)); b_ = edate(P["cod"], 12 * t)
        return max(0, min(1, (b_ - max(P["fpeaje"], a_)).days / (b_ - a_).days))
    E = {}; EV = {}; AH = {}; OP = {}; PJ = {}; EB = {}; DP = {}; PU = {}; IU = {}; TR = {}; FU = {}
    for t in T:
        E[t] = 0 if t < 1 else Pkw / 1000 * frec * Y[t - 1]
        EV[t] = 0 if t < 1 else min(E[t], P["Consumo_Anual"] * (1 + P["Crecimiento_Consumo"]) ** (t - 1) / 1000)
        AH[t] = EV[t] * 1000 * P["Tarifa_Evitable"] * pc["fT"] * (1 + pc["escT"]) ** max(0, t - 1)
        OP[t] = 0 if t < 1 else opex1 * pc["fO"] * (1 + P["Escalacion_OPEX"]) ** (t - 1)
        PJ[t] = E[t] * 1000 * pc["pj"] * (frac_peaje(t) if t >= 1 else 0)
        EB[t] = AH[t] - OP[t] - PJ[t]
        DP[t] = (dep_eq if 1 <= t <= P["Vida_Fiscal_Equipos"] else 0) + (dep_civ if 1 <= t <= min(P["Vida_Fiscal_Civil"], H) else 0)
        pu = P["Tasa_Participacion"] * (EB[t] - DP[t]) if part_on else 0
        PU[t] = pu if esc else max(0, pu)
        da = dedad if 1 <= t <= P["Vida_Fiscal_Equipos"] else 0
        iu = P["Tasa_IR"] * (EB[t] - DP[t] - PU[t] - da)
        IU[t] = iu if esc else max(0, iu)
        TR[t] = (-terr_cost if t == -1 else (resid_net if t == H else 0)) if terr else 0
        if t == -1: FU[t] = -(K + IVA) * fase + TR[t]
        elif t == 0: FU[t] = -(K + IVA) * (1 - fase) + (IVA * fase if recup else 0) + TR[t]
        else: FU[t] = EB[t] - PU[t] - IU[t] + (IVA * (1 - fase) if (t == 1 and recup) else 0) + TR[t]
    # deuda
    D = int(pc["deb"]) * pc["lev"] * (K + finT * terr * terr_cost); rd = pc["rd"]; gr = int(pc["gr"]); pl = int(pc["plazo"])
    idc = D * rd * (fase + (1 - fase) * P["IDC_Frac_Tramo0"]); Dt = D + idc; n = max(1, pl - gr)
    pmt = 0 if Dt <= 0 else (Dt / n if rd == 0 else Dt * rd / (1 - (1 + rd) ** (-n)))
    IN = {}; AM = {}; bal = Dt
    for t in T:
        if t < 1 or D == 0 or t > pl: IN[t] = 0; AM[t] = 0
        elif t <= gr:
            IN[t] = Dt * rd; AM[t] = Dt if (pl <= gr and t == pl) else 0
        else:
            IN[t] = bal * rd; AM[t] = pmt - IN[t] if pl > gr else 0; bal -= AM[t]
    PL = {}; IL = {}; CF = {}; EQ = {}; DS = []
    for t in T:
        pl_ = P["Tasa_Participacion"] * (EB[t] - DP[t] - IN[t]) if part_on else 0
        PL[t] = pl_ if esc else max(0, pl_)
        da = dedad if 1 <= t <= P["Vida_Fiscal_Equipos"] else 0
        il = P["Tasa_IR"] * (EB[t] - DP[t] - IN[t] - PL[t] - da); IL[t] = il if esc else max(0, il)
        CF[t] = 0 if t < 1 else EB[t] - PL[t] - IL[t] + (IVA * (1 - fase) if (t == 1 and recup) else 0) + TR[t]
        if t == -1: EQ[t] = FU[t] + D * fase
        elif t == 0: EQ[t] = FU[t] + D * (1 - fase)
        else: EQ[t] = CF[t] - IN[t] - AM[t]
        if IN[t] + AM[t] > 0: DS.append(CF[t] / (IN[t] + AM[t]))
    # Exergy
    UX = {}; FX = {}; GX = {}
    for t in T:
        g = (1 + P["Escalacion_OPEX"]) ** (t - 1)
        ux = 0
        if t == -1: ux = (P["Fee_Gerencia_Pct"] - P["Costo_Gerencia_Pct"]) * sub * fase
        elif t == 0: ux = (P["Fee_Gerencia_Pct"] - P["Costo_Gerencia_Pct"]) * sub * (1 - fase)
        if 1 <= t <= H:
            ux += (1 - terr) * (P["Renta_Terreno_ha"] * ha - P["Predial_Terreno"]) * g + (P["Fee_OM_kWp"] - P["Costo_OM_Exergy_kWp"]) * Pkw * g
        UX[t] = ux
        ix = -max(0, ux) * P["Tasa_Efectiva_Exergy"]
        tx = 0 if terr else (-terr_cost if t == -1 else (resid_net if t == H else 0))
        FX[t] = ux + ix + tx; GX[t] = FU[t] + FX[t]
    fu = [FU[t] for t in T]; eq = [EQ[t] for t in T]; fx = [FX[t] for t in T]; gx = [GX[t] for t in T]
    cum = np.cumsum(fu); nneg = int(np.sum(cum < 0))
    pb = (nneg - 2) + (-cum[nneg - 1]) / fu[nneg] if nneg < len(fu) else None
    df = {t: (0 if t < 1 else 1 / (1 + r) ** t) for t in T}
    lcoe = (K + sum((OP[t] + PJ[t]) * df[t] for t in T)) / sum(E[t] * df[t] for t in T)
    out = dict(TIR=irr(fu), VAN=npv0(r, fu), PB=pb, LCOE=lcoe, TIR_eq=(irr(eq) if D > 0 else None), VAN_eq=npv0(r, eq), DSCR_min=(min(DS) if (D > 0 and DS) else None), DSCR_avg=(sum(DS) / len(DS) if (D > 0 and DS) else None),
               Ahorro1=AH[1], Aporte_eq=-(EQ[-1] + EQ[0]), VAN_X=npv0(r, fx), TIR_G=irr(gx), E1=E[1], NoRec=sum(E[t] - EV[t] for t in T), Cob=E[1] * 1000 / P["Consumo_Anual"], Nominal_X=sum(fx))
    return out, dict(FU=fu, EQ=eq, FX=fx)


if __name__ == "__main__":
    VAL, OUT = sys.argv[1], sys.argv[2]
    B = Book(VAL); P = load_inputs(B)
    wsm = B.wb["Motor_Sens"]
    ncases = 0
    while wsm.cell(row=4, column=2 + ncases).value is not None:
        ncases += 1
    rows = []; worst = 0; nfail = 0; ncmp = 0
    for j in range(ncases):
        X = gl(2 + j)
        pc = {k: wsm[f"{X}{r}"].value for k, r in PARAM_ROWS.items()}
        out, _ = run_case(P, pc)
        for k, r in OUT_ROWS.items():
            book = wsm[f"{X}{r}"].value; mine = out[k]
            if isinstance(book, (int, float)) and isinstance(mine, (int, float)) and not (isinstance(mine, float) and math.isnan(mine)):
                d = abs(book - mine) / max(1.0, abs(book)); ok = d <= 1e-6; worst = max(worst, d); nfail += (not ok); ncmp += 1
                rows.append((j, wsm[f"{X}4"].value, k, book, mine, d, "●" if ok else "■"))
                if not ok and nfail <= 15: print("  ■", j, wsm[f"{X}4"].value, k, book, mine)
            else:
                rows.append((j, wsm[f"{X}4"].value, k, book, mine, None, "—"))
    with open(OUT, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(["caso", "nombre", "kpi", "libro", "sombra", "delta_rel", "estado"]); [w.writerow(x) for x in rows]
    print(f"sombra v2.0: {ncases} casos · {ncmp} KPI comparados · {nfail} difieren · peor Δrel {worst:.3g}")
    # hojas principales (caso activo)
    base = {k: wsm[f"B{r}"].value for k, r in PARAM_ROWS.items()}
    out, vec = run_case(P, base)
    ws8 = B.wb["08_Flujo"]; ws9 = B.wb["09_Exergy"]
    d8 = max(abs(ws8[f"{C(t)}{FLUJO['fcf']}"].value - vec["FU"][i]) for i, t in enumerate(TS))
    de = max(abs(ws8[f"{C(t)}{FLUJO['eq']}"].value - vec["EQ"][i]) for i, t in enumerate(TS))
    d9 = max(abs(ws9[f"{C(t)}{EXERGY['fcf']}"].value - vec["FX"][i]) for i, t in enumerate(TS))
    print(f"08_Flujo FCF Δmáx {d8:.3g} · equity Δmáx {de:.3g} · 09_Exergy FCF Δmáx {d9:.3g} · TIR_Proyecto libro {B.v('TIR_Proyecto'):.6%} sombra {out['TIR']:.6%} · VAN_Exergy {B.v('VAN_Exergy'):,.2f} / {out['VAN_X']:,.2f}")

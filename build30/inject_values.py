# -*- coding: utf-8 -*-
"""Inyecta valores calculados (de una copia recalculada por LibreOffice) como <v> cacheados en el
archivo original de openpyxl, preservando formato y gráficos. Uso: inject_values.py original.xlsx calc.xlsx out.xlsx"""
import sys, re, zipfile, shutil, datetime
from xml.sax.saxutils import escape
from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel

orig, calc, out = sys.argv[1], sys.argv[2], sys.argv[3]
wbv = load_workbook(calc, data_only=True)
values = {ws.title: {c.coordinate: c.value for row in ws.iter_rows() for c in row if c.value is not None} for ws in wbv.worksheets}

zin = zipfile.ZipFile(orig)
wbxml = zin.read("xl/workbook.xml").decode("utf8")
rels = zin.read("xl/_rels/workbook.xml.rels").decode("utf8")
rid_target = dict(re.findall(r'<Relationship[^>]*Id="(rId\d+)"[^>]*Target="([^"]+)"', rels))
rid_target.update({a: b for b, a in re.findall(r'<Relationship[^>]*Target="([^"]+)"[^>]*Id="(rId\d+)"', rels)})
sheets = re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="(rId\d+)"', wbxml)
title_by_path = {}
for nm, rid in sheets:
    tgt = rid_target[rid]
    path = "xl/" + tgt if not tgt.startswith("/") else tgt[1:]
    title_by_path[path] = nm.replace("&amp;", "&")

pat = re.compile(r'<c r="([A-Z]+\d+)"((?: [a-z]+="[^"]*")*)><f>(.*?)</f><v></v></c>', re.S)
stats = {"num": 0, "str": 0, "none": 0}


def make_sub(vals):
    def sub(m):
        coord, attrs, f = m.group(1), m.group(2), m.group(3)
        v = vals.get(coord)
        if v is None:
            stats["none"] += 1
            return m.group(0)
        if isinstance(v, bool):
            stats["num"] += 1
            return f'<c r="{coord}"{attrs} t="b"><f>{f}</f><v>{1 if v else 0}</v></c>'
        if isinstance(v, (int, float)):
            stats["num"] += 1
            return f'<c r="{coord}"{attrs}><f>{f}</f><v>{repr(float(v)) if isinstance(v, float) else v}</v></c>'
        if isinstance(v, (datetime.datetime, datetime.date)):
            stats["num"] += 1
            return f'<c r="{coord}"{attrs}><f>{f}</f><v>{to_excel(v)}</v></c>'
        s = str(v)
        if s.startswith("#"):
            stats["none"] += 1
            return m.group(0)
        stats["str"] += 1
        return f'<c r="{coord}"{attrs} t="str"><f>{f}</f><v>{escape(s)}</v></c>'
    return sub


zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename in title_by_path:
        title = title_by_path[item.filename]
        xml = data.decode("utf8")
        xml = pat.sub(make_sub(values.get(title, {})), xml)
        data = xml.encode("utf8")
    zout.writestr(item, data)
zout.close()
print("inyectados:", stats)

# -*- coding: utf-8 -*-
"""extract_model.py RAW.xlsx CALC.xlsx OUT_DIR — v0 (F1 del artefacto, doc 17): vuelca el libro a JSON para el motor JS y la maqueta.
  names.json  : todos los nombres definidos → hoja, referencia, valor(es) calculados (LibreOffice/Excel) y fórmula(s)
  motor.json  : Motor_Sens completo por caso (parámetros, escalares, salidas y 28 bloques anuales) con etiquetas de fila y fórmulas de parámetros
  sheets.json : celdas no vacías (valor + fórmula) de todas las hojas salvo el Motor (textos legales, tablas, layouts)
  meta.json   : versión, fecha, SHA-256 de los dos libros, conteos"""
import sys, json, hashlib, datetime as dt
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

raw_path, calc_path, out = sys.argv[1], sys.argv[2], sys.argv[3]
import os; os.makedirs(out, exist_ok=True)
RAW = load_workbook(raw_path)                    # fórmulas
CAL = load_workbook(calc_path, data_only=True)   # valores


def js(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat()[:10]
    return v


def cells_of(ref, wb):
    sh, rg = ref.replace("'", "").replace("$", "").split("!")
    ws = wb[sh]
    if ":" in rg:
        return sh, rg, [[js(c.value) for c in row] for row in ws[rg]]
    return sh, rg, js(ws[rg].value)


names = {}
for nm, dn in RAW.defined_names.items():
    try:
        sh, rg, vals = cells_of(dn.attr_text, CAL)
        _, _, forms = cells_of(dn.attr_text, RAW)
    except Exception as e:
        names[nm] = {"error": str(e), "ref": dn.attr_text}; continue
    names[nm] = {"sheet": sh, "ref": rg, "value": vals, "formula": forms}
json.dump(names, open(f"{out}/names.json", "w"), ensure_ascii=False, indent=0)

# ---- Motor por etiqueta
wm, wmf = CAL["Motor_Sens"], RAW["Motor_Sens"]
labels = {"params": [], "scalars": [], "outputs": [], "blocks": []}
rows = {}
block = None; block_rows = {}
for r in range(5, wm.max_row + 1):
    a = wm.cell(r, 1).value
    if a is None:
        continue
    if isinstance(a, str) and a.startswith("Bloque "):
        block = a.split(" ")[1]; block_rows[block] = {}; labels["blocks"].append(block); continue
    if block is not None and isinstance(a, (int, float)) and not isinstance(a, bool):
        block_rows[block][int(a)] = r
    elif isinstance(a, str):
        rows[a] = r
# clasificación por posición (parámetros hasta la primera fila de escalares 'Factor de recorte'; salidas desde 'TIR proyecto')
order = sorted(rows.items(), key=lambda kv: kv[1])
r_scal0 = rows["Factor de recorte"]; r_out0 = rows["TIR proyecto"]
for lab, r in order:
    (labels["params"] if r < r_scal0 else labels["scalars"] if r < r_out0 else labels["outputs"]).append(lab)
cases = []
for c in range(2, wm.max_column + 1):
    name = wm.cell(4, c).value
    if not name:
        continue
    col = get_column_letter(c)
    cs = {"name": name, "col": col, "params": {}, "paramFormulas": {}, "scalars": {}, "outputs": {}, "blocks": {}}
    for lab in labels["params"]:
        cs["params"][lab] = js(wm.cell(rows[lab], c).value); cs["paramFormulas"][lab] = wmf.cell(rows[lab], c).value
    for lab in labels["scalars"]:
        cs["scalars"][lab] = js(wm.cell(rows[lab], c).value)
    for lab in labels["outputs"]:
        cs["outputs"][lab] = js(wm.cell(rows[lab], c).value)
    for b, br in block_rows.items():
        cs["blocks"][b] = [js(wm.cell(br[t], c).value) for t in sorted(br)]
    cases.append(cs)
ts = sorted(next(iter(block_rows.values())).keys())
# fórmulas del caso Custom (columna B) para escalares, salidas y bloques: la referencia de la transliteración
custom_formulas = {"scalars": {lab: wmf.cell(rows[lab], 2).value for lab in labels["scalars"]},
                   "outputs": {lab: wmf.cell(rows[lab], 2).value for lab in labels["outputs"]},
                   "blocks": {b: {str(t): wmf.cell(br[t], 2).value for t in sorted(br)} for b, br in block_rows.items()}}
json.dump({"labels": labels, "t": ts, "cases": cases, "customFormulas": custom_formulas}, open(f"{out}/motor.json", "w"), ensure_ascii=False, indent=0)

# ---- hojas (valor + fórmula), sin el Motor
sheets = {}
for ws in RAW.worksheets:
    if ws.title == "Motor_Sens":
        continue
    wsv = CAL[ws.title]
    cells = {}
    for row in ws.iter_rows():
        for cf in row:
            if cf.value is None:
                continue
            cv = wsv.cell(cf.row, cf.column).value
            entry = {"v": js(cv)}
            if isinstance(cf.value, str) and cf.value.startswith("="):
                entry["f"] = cf.value
            cells[cf.coordinate] = entry
    sheets[ws.title] = {"state": ws.sheet_state, "cells": cells, "merged": [str(m) for m in ws.merged_cells.ranges]}
json.dump(sheets, open(f"{out}/sheets.json", "w"), ensure_ascii=False, indent=0)

sha = lambda p: hashlib.sha256(open(p, "rb").read()).hexdigest()
meta = {"version": names.get("Version", {}).get("value"), "fecha_analisis": names.get("Fecha_Analisis", {}).get("value"),
        "raw": {"path": raw_path, "sha256": sha(raw_path)}, "calc": {"path": calc_path, "sha256": sha(calc_path)},
        "n_names": len(names), "n_cases": len(cases), "n_params": len(labels["params"]), "n_scalars": len(labels["scalars"]), "n_outputs": len(labels["outputs"]), "n_blocks": len(labels["blocks"]),
        "extracted": dt.datetime.utcnow().isoformat() + "Z"}
json.dump(meta, open(f"{out}/meta.json", "w"), ensure_ascii=False, indent=1)
print(json.dumps(meta, ensure_ascii=False, indent=1))

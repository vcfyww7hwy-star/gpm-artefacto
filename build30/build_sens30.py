# -*- coding: utf-8 -*-
"""10_Sensibilidad v2.0 — retícula y páginas de la v1.3 con la nueva arquitectura de casos:
· §A los cuatro casos (Custom · Conservador · Base · Favorable) leídos del Motor, con su definición viva (bloque B de 01) y dos gráficos;
· §A.2 P50 vs P90 del Custom (nombres P50_* / P90_*); §B–§E tornado, matrices, deuda y piso sobre el Custom;
· §F barrido de potencia 5–8 MWp (lista de 01) y ratio DC/AC; §G terreno; §H deuda máxima. Desaparece §I (estaba en §A).
· el Custom se marca con ◆ y borde ciruela (formato condicional); los cuatro casos llevan su color sólo en cabeceras, series y ◆."""
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.chart import Reference
from openpyxl.worksheet.pagebreak import Break
from xl_helpers import *
from build_core import (S10, SM, C, col, NSHEETS, OUT_ROWS, PARAM_ROWS, SCAL_ROWS)
from build_core import CASE_KEYS
from build_content import (N_CASES, CASE_BASE, CASE_X, CASE_C, CASE_B, CASE_F, CASE_ESC, CASE_P50, CASE_P90, CASE_PISO, TORNADO, MAT_START, EQ_START, LEV_START, CASE_TERR_ALT,
                           N_SWEEP_P, N_SWEEP_R, N_SWEEP_LEV, SWEEP_P_START, SWEEP_R_START,
                           SWEEP_A_START, DM_START, PRECIO_STEPS, TS_START, TX_START, SENS_ANCHORS, TORNADO_SHORT_COL, TORNADO_SHORT,
                           T_REP, BR0, BR1, BR2, BR3, BR4, BR5, TORNADO_RANK_COLS, mo, mcol, cf_text, cf_scale)

LC = 12          # última columna impresa (L)
RH = 15          # alto de fila de tabla
ACT_SIDE = Side(style="thin", color=X_COL)   # v2.0: el Custom se marca en ciruela


def cf_active(ws, rng_, formula, sides=("left", "right")):
    """Borde ciruela (formato condicional) en las celdas del Custom: la fórmula se evalúa con referencias relativas a la primera celda del rango."""
    kw = {s: ACT_SIDE for s in sides}
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[formula], border=Border(**kw), font=Font(bold=True, color=CARBON)))


def hdr_cases(ws, r, c1, texts, height=30):
    """Cabecera de casos (columnas): texto o fórmula por columna, negrita grafito, regla carbón; la primera celda (c1) es la etiqueta de la fila."""
    hdr(ws, r, 2, c1 + len(texts) - 1, [""] * (c1 - 2) + texts, height=height)
    for cc in range(c1, c1 + len(texts)):
        ws.cell(row=r, column=cc).alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)


def build_sensibilidad(wb):
    ws = wb.create_sheet(S10)
    widths(ws, {"A": 2, "B": 40})
    for cc in range(3, LC + 1):
        ws.column_dimensions[col(cc)].width = 11
    sheet_header(ws, "10 · Sensibilidad", f"Qué mueve el resultado y cuánto instalar, en vivo desde el Motor ({N_CASES} casos): cuatro casos y puente (§A), tornado (§B), matrices (§C), deuda (§D, §H), piso (§E), potencia (§F), terreno y reemplazo (§G). ◆ = Custom.", 10, last_col=LC, total=NSHEETS)
    ws.row_dimensions[4].height = 10
    breaks = []
    r = 5
    # ================================================================ parámetros (dos por fila)
    section(ws, r, 2, LC, "Parámetros de sensibilidad (editables)", guide="pasos del tornado, las matrices y los barridos; el resto de la hoja se recalcula", guide_col=6)
    params = [("Sens_CAPEX", "Variación del CAPEX (±)", 0.15, FMT_PCT), ("Sens_Tarifa", "Variación de la tarifa evitable (±)", 0.15, FMT_PCT),
              ("Sens_OPEX_Up", "OPEX al alza (+)", 0.30, FMT_PCT), ("Sens_OPEX_Dn", "OPEX a la baja (−)", 0.15, FMT_PCT),
              ("Sens_Peaje", "Peaje SGDA ($/kWh inyectado, desde 2029)", 0.015, FMT_KWH), ("Sens_EscTarifa", "Escalación de la tarifa: aumento frente al Custom (pp/año)", 0.01, FMT_PCT),
              ("Sens_Disponibilidad", "Disponibilidad: reducción frente al Custom (pp)", 0.02, FMT_PCT), ("Sens_EscCAPEX", "Escalación del CAPEX: aumento (pp/año)", 0.02, FMT_PCT),
              ("Sens_Peaje_kW", "Peaje por potencia ($/kW-mes sobre la AC)", 0.5, FMT_DEC2), ("Sweep_AC_Fija", "Potencia AC fija del barrido §F.3 (kWac)", 3800, FMT_INT)]
    r += 1
    for i in range(0, len(params), 2):
        for j, (nm, lab, v, fmt) in enumerate(params[i:i + 2]):
            cl, cv, cn = (2, 3, 4) if j == 0 else (6, 10, 11)   # izquierda: B · C · D  |  derecha: F:I (etiqueta) · J · K
            if j == 1:
                lc_ = ws.cell(row=r, column=cl, value=lab); lc_.font = Font(name=FONT, size=SZ_BODY, color=CARBON); lc_.alignment = Alignment(vertical="center")
                ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=9)
                for cc in range(6, 10): ws.cell(row=r, column=cc).border = B_BOTTOM
            else:
                label(ws, r, cl, lab)
            inp(ws, r, cv, v, fmt=fmt)
            c = ws.cell(row=r, column=cn, value=nm); c.font = font(size=8.5, color=GRAFITO); c.border = B_BOTTOM
            name(wb, nm, S10, f"${col(cv)}${r}")
        for cc in (5, 12): ws.cell(row=r, column=cc).border = B_BOTTOM
        ws.row_dimensions[r].height = RH
        r += 1
    lists = [("Pasos matriz CAPEX (Δ sobre el Custom)", [-0.15, -0.075, 0, 0.075, 0.15], FMT_SIGNPCT, "Mat_CAPEX"),
             ("Pasos matriz tarifa (Δ sobre la tarifa evitable)", [-0.15, -0.075, 0, 0.075, 0.15], FMT_SIGNPCT, "Mat_Tarifa"),
             ("Tasas de deuda a evaluar", [0.075, 0.0825, 0.09, 0.10, 0.11], FMT_PCT2, "Sens_Tasas"),
             ("Plazos de deuda a evaluar (años, incl. gracia)", [8, 10, 12], "0", "Sens_Plazos"),
             ("Apalancamientos a evaluar (deuda / CAPEX)", [0.5, 0.6, 0.7, 0.8], FMT_PCT0, "Sens_Lev"),
             ("Barrido de potencia DC (kWp) §F.1 — la lista de 01", [5000, 6000, 7000, 8000], FMT_INT, "Sweep_P"),
             ("Barrido de ratio DC/AC §F.2 y §F.3", [1.10, 1.20, 1.32, 1.40, 1.50], "0.00", "Sweep_Ratio"),
             ("Apalancamientos para la deuda máxima §H", [0.4, 0.5, 0.6, 0.7, 0.8], FMT_PCT0, "Sweep_Lev")]
    for lab, vals, fmt, nm in lists:
        label(ws, r, 2, lab)
        for i, v in enumerate(vals):
            inp(ws, r, 3 + i, v, fmt=fmt)
        for cc in range(3 + len(vals), LC + 1): ws.cell(row=r, column=cc).border = B_BOTTOM
        name(wb, nm, S10, f"$C${r}:${col(2+len(vals))}${r}"); ws.row_dimensions[r].height = RH
        r += 1
    ws.row_dimensions[r].height = 8; r += 1
    # ================================================================ A · Los cuatro casos
    section(ws, r, 2, LC, "A · Los cuatro casos — Custom · Conservador · Base · Favorable", guide="leídos del Motor con la potencia, el contrato, el terreno y la deuda vigentes; el Custom gobierna 04–09 y la portada", guide_col=6)
    SENS_ANCHORS["A"] = f"B{r}"
    r += 1
    hdr(ws, r, 2, 7, ["Indicador", "", "", "", "", "Δ Custom − Base"], height=30)
    for j, key in enumerate(CASE_KEYS):
        caso_hdr(ws, r, 3 + j, key, size=SZ_TABLE)
    r += 1
    a0 = r
    A_ROWS = [("Energía año 1 [MWh]", "E1", FMT_INT), ("Ahorro año 1 [USD]", "Ahorro1", FMT_USD), ("TIR del proyecto (sin deuda)", "TIR", FMT_PCT2), ("VAN @ tasa de descuento [USD]", "VAN", FMT_USD),
              ("Payback desde COD [años]", "PB", FMT_YRS), ("LCOE [$/MWh]", "LCOE", FMT_DEC1), ("TIR del accionista (con deuda)", "TIR_eq", FMT_PCT2), ('="VAN del accionista @ "&TEXT(Tasa_Descuento_Equity,"0%")&" [USD]"', "VAN_eq", FMT_USD),
              ("DSCR mínimo", "DSCR_min", FMT_X), ("Aporte de capital con deuda [USD]", "Aporte_eq", FMT_USD), ("VAN negocio Exergy [USD]", "VAN_X", FMT_USD), ("TIR consolidada del grupo", "TIR_G", FMT_PCT2)]
    for lab, key, fmt in A_ROWS:
        label(ws, r, 2, lab, bold=(key in ("TIR", "VAN", "TIR_eq")))
        for j, ck in enumerate(CASE_KEYS):
            caso_val(ws, r, 3 + j, f"={mo(CASE_ESC[ck], key)}", fmt=fmt, bold=(ck == "X"), size=SZ_BODY)
        calc(ws, r, 7, f'=IFERROR(C{r}-E{r},"")', fmt=fmt, color=GRAFITO)
        for cc in range(8, LC + 1): ws.cell(row=r, column=cc).border = B_BOTTOM
        ws.row_dimensions[r].height = RH
        r += 1
    for j, key in enumerate(["E1", "Ahorro1", "TIR", "VAN", "PB", "LCOE", "TIReq", "VANeq", "DSCR", "Aporte", "VANX", "TIRG"]):
        for k, tag in enumerate(CASE_KEYS):
            name(wb, f"{tag}_{key}", S10, f"${col(3+k)}${a0+j}")
    cf_scale(ws, f"C{a0+2}:F{a0+2}"); cf_scale(ws, f"C{a0+3}:F{a0+3}", mid_num=0); cf_scale(ws, f"C{a0+6}:F{a0+6}"); cf_text(ws, f"C{a0+8}:F{a0+8}", 1, "DSCR_Objetivo")
    a_last = r - 1
    # definición viva de cada caso (bloque B de 01) — una línea por caso, texto del color del caso
    def esc_def(k):
        e = lambda nm: f"INDEX({nm},{k})"
        return (f'="{CASO_NOMBRE[CASE_KEYS[k-1]]} = "&{e("Esc_Energia")}&" · CAPEX "&IF(N({e("Esc_CAPEX_Fijo_Wp")})>0,"fijo "&TEXT({e("Esc_CAPEX_Fijo_Wp")},"0.00")&" $/Wp","bottom-up × "&TEXT({e("Esc_Factor_CAPEX")},"0.00"))'
                f'&IF({e("Esc_Escalacion_CAPEX")}>0," +"&TEXT({e("Esc_Escalacion_CAPEX")},"0%")&"/año","")&" · OPEX × "&TEXT({e("Esc_Factor_OPEX")},"0.00")&" · peaje "&TEXT({e("Esc_Peaje")}*100,"0.0")&" ¢/kWh desde "&TEXT(Fecha_Peaje,"yyyy")&" · tarifa "&IF({e("Esc_EscTarifa")}=0,"plana","+"&TEXT({e("Esc_EscTarifa")},"0.0%")&"/año")&" · disp. "&TEXT({e("Esc_Disponibilidad")},"0%")'
                + ('&" · resto: capa de diseño de 01"' if k == 1 else ''))
    for k, ck in enumerate(CASE_KEYS):
        c = ws.cell(row=r, column=2, value=esc_def(k + 1)); c.font = Font(name=FONT, size=SZ_TABLE, color=CASO_COL[ck], bold=True); c.alignment = Alignment(vertical="center")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=LC); ws.row_dimensions[r].height = RH
        r += 1
    chip(ws, r, 2, '=Estado_Custom&"  ·  "&Estado_Entregado', kind="custom", c2=LC); ws.row_dimensions[r].height = RH
    r += 1
    ws.row_dimensions[r].height = 8; r += 1
    breaks.append(r - 1)
    # ================================================================ A.3 · Puente v2.0 → v3.0 (tabla + cascada) — v3.0 (V6)
    section(ws, r, 2, LC, "A.3 · Puente v2.0 → v3.0 — del Base con la ronda 2 en neutro al Base actual, un parámetro a la vez", guide="cada escalón enciende un parámetro con los anteriores ya activos; el último ≡ Base (control H7)", guide_col=7)
    SENS_ANCHORS["A3"] = f"B{r}"
    r += 1
    hdr(ws, r, 2, LC, ["Escalón", "TIR proyecto", "Δ (pp)", "VAN proyecto", "TIR accionista", "Δ (pp)", "DSCR mín", "VAN accionista", "VAN Exergy", "Payback", "LCOE"], height=30)
    r += 1
    br0 = r
    BR_ROWS = [(BR0, "Definición v2.0 (ronda 2 en neutro)"),
               (BR1, '="+ escalación del CAPEX "&TEXT(INDEX(Esc_Escalacion_CAPEX,3),"0%")&"/año desde "&TEXT(Fecha_Precios,"mmm-yyyy")'),
               (BR2, '="+ disponibilidad "&TEXT(INDEX(Esc_Disponibilidad,3),"0%")'),
               (BR3, '="+ reemplazo de inversores en t = "&Reemplazo_Anio&" ("&Reemplazo_Pagador&")"'),
               (BR4, '="+ construcción de "&Meses_Construccion&" meses (IDC)"'),
               (BR5, '="+ tasa del accionista "&TEXT(Tasa_Descuento_Equity,"0%")&" y resto de la capa de diseño ≡ Base"')]
    for k, (ci, lab) in enumerate(BR_ROWS):
        label(ws, r, 2, lab, bold=(k in (0, 5)))
        calc(ws, r, 3, f"={mo(ci,'TIR')}", fmt=FMT_PCT2, bold=(k in (0, 5)))
        calc(ws, r, 4, ("=0" if k == 0 else f'=IFERROR((C{r}-C{r-1})*100,"")'), fmt=FMT_PP, color=GRAFITO)
        calc(ws, r, 5, f"={mo(ci,'VAN')}", fmt=FMT_USD)
        calc(ws, r, 6, f"={mo(ci,'TIR_eq')}", fmt=FMT_PCT2)
        calc(ws, r, 7, ("=0" if k == 0 else f'=IFERROR((F{r}-F{r-1})*100,"")'), fmt=FMT_PP, color=GRAFITO)
        calc(ws, r, 8, f"={mo(ci,'DSCR_min')}", fmt=FMT_X); calc(ws, r, 9, f"={mo(ci,'VAN_eq')}", fmt=FMT_USD); calc(ws, r, 10, f"={mo(ci,'VAN_X')}", fmt=FMT_USD)
        calc(ws, r, 11, f"={mo(ci,'PB')}", fmt=FMT_YRS); calc(ws, r, 12, f"={mo(ci,'LCOE')}", fmt=FMT_DEC1)
        ws.row_dimensions[r].height = RH
        r += 1
    br5 = r - 1
    cf_active(ws, f"B{br5}:{col(LC)}{br5}", "TRUE", sides=("top", "bottom"))
    callout(ws, r, 2, LC, '="Lectura: la TIR del proyecto pasa de "&TEXT(C' + str(br0) + ',"0.00%")&" (definición v2.0) a "&TEXT(C' + str(br5) + ',"0.00%")&" (Base actual); la del accionista, de "&TEXT(F' + str(br0) + ',"0.00%")&" a "&TEXT(F' + str(br5) + ',"0.00%")&" (los meses de construcción y la tasa del accionista sólo mueven al accionista). Los valores de los escalones son supuestos «por confirmar» del bloque B y de la capa de diseño de 01. "&Estado_Neutro', height=30)
    r += 1
    ws.row_dimensions[r].height = 6; r += 1
    # cascada (TIR del proyecto): auxiliares en U:X — base invisible + Δ coloreado; totales en los extremos
    ax = r
    ws.cell(row=ax - 1, column=21, value="Auxiliares de la cascada §A.3 (no editar): categoría · base · altura").font = font(size=8.5, color=GRAFITO)
    casc_rows = []
    for k in range(6):
        rr = ax + k
        rb = br0 + k
        if k == 0:
            cat = f'="v2.0  "&TEXT(C{rb},"0.00%")'; base = "=0"; height = f"=C{rb}"
        elif k == 5:
            cat = f'="Base actual  "&TEXT(C{rb},"0.00%")'; base = "=0"; height = f"=C{rb}"
        else:
            cat = f'=B{rb}&"  "&TEXT(D{rb},"+0.00;−0.00")&" pp"'; base = f"=MIN(C{rb},C{rb-1})"; height = f"=ABS(C{rb}-C{rb-1})"
        for cc, f in ((21, cat), (22, base), (23, height)):
            c = ws.cell(row=rr, column=cc, value=f); c.font = font(size=8.5, color=GRAFITO); c.number_format = FMT_PCT2
        casc_rows.append(rr)
    cats_c = Reference(ws, min_col=21, min_row=ax, max_row=ax + 5)
    chart_cascade(ws, f"B{r}", "A.3 · Puente de la TIR del proyecto: libro anterior → Base actual (un parámetro por escalón)", cats_c,
                  Reference(ws, min_col=22, min_row=ax, max_row=ax + 5), Reference(ws, min_col=23, min_row=ax, max_row=ax + 5),
                  [GRAFITO, TERRACOTA, TERRACOTA, TERRACOTA, TERRACOTA, GRAFITO], w=19.0, h=6.2, from_rows=False, y_min=0.07, label_idxs=(0, 5))
    NCHC = 14
    for k in range(r, r + NCHC):
        ws.row_dimensions[k].height = RH          # 14 × 14 = 196 pt ≥ 6,2 cm (176 pt)
    r += NCHC
    ws.row_dimensions[r].height = 6; r += 1
    # ================================================================ A.2 · P50 vs P90 del Custom (tabla compacta)
    section(ws, r, 2, LC, "A.2 · Custom con P50 y con P90 (energía)", guide="◆ = escenario de energía del Custom; el resto de entradas como en 01", guide_col=6)
    SENS_ANCHORS["A2"] = f"B{r}"
    r += 1
    hdr(ws, r, 2, 5, ["Indicador", '=IF(Eff_Scen=1,"◆ P50","P50")', '=IF(Eff_Scen=2,"◆ P90","P90")', "Δ (P90 − P50)"], height=18)
    rows = [("TIR del proyecto (sin deuda)", "TIR", FMT_PCT2), ("VAN del proyecto @ tasa de descuento", "VAN", FMT_USD), ("Payback simple (años desde COD)", "PB", FMT_YRS),
            ("LCOE ($/MWh)", "LCOE", FMT_DEC1), ("TIR del accionista (con deuda)", "TIR_eq", FMT_PCT2), ("VAN del accionista", "VAN_eq", FMT_USD),
            ("DSCR mínimo", "DSCR_min", FMT_X), ("Ahorro año 1", "Ahorro1", FMT_USD)]
    r += 1
    pp_row = r
    for lab, key, fmt in rows:
        label(ws, r, 2, lab); calc(ws, r, 3, f"={mo(CASE_P50, key)}", fmt=fmt); calc(ws, r, 4, f"={mo(CASE_P90, key)}", fmt=fmt); calc(ws, r, 5, f'=IFERROR(D{r}-C{r},"")', fmt=fmt, color=GRAFITO)
        ws.row_dimensions[r].height = RH
        r += 1
    for i, key in enumerate(["TIR", "VAN", "PB", "LCOE", "TIReq", "VANeq", "DSCR", "Ahorro1"]):
        name(wb, f"P50_{key}", S10, f"$C${pp_row+i}"); name(wb, f"P90_{key}", S10, f"$D${pp_row+i}")
    cf_active(ws, f"C{pp_row}:C{pp_row+7}", "Eff_Scen=1"); cf_active(ws, f"D{pp_row}:D{pp_row+7}", "Eff_Scen=2")
    ws.row_dimensions[r].height = 8; r += 1
    breaks.append(r - 1)
    # ================================================================ gráficos de los cuatro casos (§A) + tabla del tornado (§B)
    gr = r
    catsA = Reference(ws, min_col=3, max_col=6, min_row=a0 - 1)
    # referencia del gráfico (tasa de descuento × 4) en U:X, fuera del área de impresión
    ws.cell(row=a0 - 1, column=21, value="Tasa de descuento (referencia del gráfico §A; no editar)").font = font(size=SZ_NOTE, color=GRAFITO)
    for cc in range(21, 25):
        c = ws.cell(row=a0, column=cc, value="=Tasa_Descuento"); c.font = font(size=SZ_NOTE, color=GRAFITO); c.number_format = FMT_PCT2
    caso_cols = [CASO_COL[k] for k in CASE_KEYS]
    chart_cols(ws, f"B{gr}", "A · TIR del proyecto por caso (referencia: tasa de descuento)", catsA,
               [dict(ref=Reference(ws, min_col=3, max_col=6, min_row=a0 + 2), name="TIR proyecto", color=GRAFITO, point_colors=caso_cols, labels="all", label_fmt="0.0%", label_pos="outEnd")],
               w=11.8, h=7.0, y_fmt="0%", legend=None, from_rows=True, gap=80,
               ref=dict(ref=Reference(ws, min_col=21, max_col=24, min_row=a0), name="tasa de descuento", color=ARCILLA, width=1.5, dash="dash"))
    chart_cols(ws, f"G{gr}", "A · VAN del proyecto y de Exergy por caso (miles de USD)", catsA,
               [dict(ref=Reference(ws, min_col=3, max_col=6, min_row=a0 + 3), name="VAN proyecto", color=GRAFITO, point_colors=caso_cols, labels="all", label_fmt=FMT_K, label_pos="outEnd"),
                dict(ref=Reference(ws, min_col=3, max_col=6, min_row=a0 + 10), name="VAN Exergy", color=NIEBLA, labels="all", label_fmt=FMT_K, label_pos="outEnd", label_color=PIEDRA)],
               w=11.6, h=7.0, y_fmt=FMT_K, legend="b", from_rows=True, gap=60)
    NCHA = 15
    for k in range(gr, gr + NCHA):
        ws.row_dimensions[k].height = RH          # 15 × 14 = 210 pt ≥ 7,0 cm (198 pt)
    r = gr + NCHA
    ws.row_dimensions[r].height = 6; r += 1
    # ---- §B · tabla del tornado (15 barras en v3.1; el gráfico, ordenado por amplitud, va en la página siguiente)
    section(ws, r, 2, LC, "B · Tornado — TIR del proyecto sin deuda (una variable a la vez, sobre el Custom)", guide=f"el gráfico ordena las {len(TORNADO)} barras por amplitud; la nota (n) explica cada una", guide_col=7)
    SENS_ANCHORS["B"] = f"B{r}"
    r += 1
    hdr(ws, r, 2, 10, ["Variable", "TIR caso bajo", "TIR caso alto", "Δ bajo (pp)", "Δ alto (pp)", "VAN caso bajo", "VAN caso alto", "Amplitud (pp)", "nota"], height=30)
    r += 1
    label(ws, r, 2, "TIR del Custom (referencia)", bold=True); calc(ws, r, 3, f"={mo(CASE_BASE,'TIR')}", fmt=FMT_PCT2, bold=True); calc(ws, r, 7, f"={mo(CASE_BASE,'VAN')}", fmt=FMT_USD, bold=True)
    for cc in (4, 5, 6, 8, 9, 10): ws.cell(row=r, column=cc).border = B_BOTTOM
    ws.row_dimensions[r].height = RH
    base_r = r
    r += 1
    t0 = r
    for k, (lab, lo, hi, nt) in enumerate(TORNADO):
        label(ws, r, 2, lab)
        calc(ws, r, 3, f"={mo(lo,'TIR')}", fmt=FMT_PCT2)
        if hi is not None:
            calc(ws, r, 4, f"={mo(hi,'TIR')}", fmt=FMT_PCT2); calc(ws, r, 8, f"={mo(hi,'VAN')}", fmt=FMT_USD)
        else:
            calc(ws, r, 4, f"=$C${base_r}", fmt=FMT_PCT2, color=GRAFITO); calc(ws, r, 8, f"=$G${base_r}", fmt=FMT_USD, color=GRAFITO)
        calc(ws, r, 5, f'=IFERROR((C{r}-$C${base_r})*100,0)', fmt=FMT_PP); calc(ws, r, 6, f'=IFERROR((D{r}-$C${base_r})*100,0)', fmt=FMT_PP)
        calc(ws, r, 7, f"={mo(lo,'VAN')}", fmt=FMT_USD); calc(ws, r, 9, f"=ABS(F{r}-E{r})", fmt="0.00", bold=True)
        c = ws.cell(row=r, column=10, value=f"({k+1})"); c.font = font(size=8.5, color=GRAFITO); c.alignment = Alignment(horizontal="center"); c.border = B_BOTTOM
        ws.row_dimensions[r].height = RH
        r += 1
    t1 = r - 1
    NT = t1 - t0 + 1
    # etiquetas cortas (fijas) en U y bloque de ranking por amplitud en V:Z (fuera del área de impresión): el gráfico y la portada leen el ranking
    RK_AMP, RK_IDX, RK_LAB, RK_LO, RK_HI = TORNADO_RANK_COLS   # V amplitud con desempate · W índice · X etiqueta · Y Δ bajo · Z Δ alto
    ws.column_dimensions[col(TORNADO_SHORT_COL)].width = 30; ws.column_dimensions[col(RK_LAB)].width = 30
    c = ws.cell(row=t0 - 1, column=TORNADO_SHORT_COL, value="Etiquetas cortas (orden de la tabla)"); c.font = font(size=8.5, color=GRAFITO)
    for cc, txt in ((RK_AMP, "Amplitud (desempate)"), (RK_IDX, "Ranking: fila"), (RK_LAB, "Ranking: etiqueta"), (RK_LO, "Ranking: Δ bajo"), (RK_HI, "Ranking: Δ alto")):
        c = ws.cell(row=t0 - 1, column=cc, value=txt); c.font = font(size=8.5, color=GRAFITO)
    amp_rng = f"${col(RK_AMP)}${t0}:${col(RK_AMP)}${t1}"
    for k, f in enumerate(TORNADO_SHORT):
        rr = t0 + k
        c = ws.cell(row=rr, column=TORNADO_SHORT_COL, value=f); c.font = font(size=8.5, color=GRAFITO)
        c = ws.cell(row=rr, column=RK_AMP, value=f"=I{rr}+({NT}-{k})*0.000000001"); c.font = font(size=8.5, color=GRAFITO); c.number_format = "0.00"
        c = ws.cell(row=rr, column=RK_IDX, value=f"=MATCH(LARGE({amp_rng},{k+1}),{amp_rng},0)"); c.font = font(size=8.5, color=GRAFITO); c.number_format = "0"
        c = ws.cell(row=rr, column=RK_LAB, value=f"=INDEX(${col(TORNADO_SHORT_COL)}${t0}:${col(TORNADO_SHORT_COL)}${t1},{col(RK_IDX)}{rr})"); c.font = font(size=8.5, color=GRAFITO)
        c = ws.cell(row=rr, column=RK_LO, value=f"=INDEX($E${t0}:$E${t1},{col(RK_IDX)}{rr})"); c.font = font(size=8.5, color=GRAFITO); c.number_format = FMT_PP
        c = ws.cell(row=rr, column=RK_HI, value=f"=INDEX($F${t0}:$F${t1},{col(RK_IDX)}{rr})"); c.font = font(size=8.5, color=GRAFITO); c.number_format = FMT_PP
    ws.row_dimensions[r].height = 8; r += 1
    breaks.append(r - 1)
    # ---- §B · gráfico (barras ordenadas por amplitud) + notas
    gr = r
    chart_tornado(ws, f"B{gr}", "Tornado · Δ TIR del proyecto en puntos porcentuales, ordenado por amplitud (caso bajo niebla · caso alto terracota)", Reference(ws, min_col=RK_LAB, min_row=t0, max_row=t1),
                  Reference(ws, min_col=RK_LO, min_row=t0, max_row=t1), Reference(ws, min_col=RK_HI, min_row=t0, max_row=t1), w=19.0, h=9.6)
    NCH = 21
    for k in range(gr, gr + NCH):
        ws.row_dimensions[k].height = RH          # 21 × (15 − 1) = 294 pt ≥ 9,6 cm (272 pt)
    r = gr + NCH
    ws.row_dimensions[r].height = 6; r += 1
    for k, (lab, lo, hi, nt) in enumerate(TORNADO):
        c = ws.cell(row=r, column=2, value=(nt if (nt.startswith("=") or nt.startswith("(")) else f"({k+1}) {nt}"))
        c.font = Font(name=FONT, size=8.5, color=GRAFITO); c.alignment = Alignment(vertical="center"); ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=LC); ws.row_dimensions[r].height = 13
        r += 1
    ws.row_dimensions[r].height = 8; r += 1
    breaks.append(r - 1)
    # ================================================================ C · Matriz
    section(ws, r, 2, LC, "C · Matriz TIR y VAN del proyecto (sin deuda): Δ CAPEX (filas) × Δ tarifa evitable (columnas)", guide="la celda con borde es el Custom (0 %, 0 %)", guide_col=7)
    SENS_ANCHORS["C"] = f"B{r}"
    r += 1
    hdr(ws, r, 2, 7, ["TIR proyecto  ·  Δ CAPEX ↓  Δ tarifa →"] + [f"=INDEX(Mat_Tarifa,{j+1})" for j in range(5)], height=18)
    for j in range(5): ws.cell(row=r, column=3 + j).number_format = FMT_SIGNPCT
    h1 = r
    r += 1
    m0 = r
    for i in range(5):
        calc(ws, r, 2, f"=INDEX(Mat_CAPEX,{i+1})", fmt=FMT_SIGNPCT, bold=True, align="left")
        for j in range(5):
            calc(ws, r, 3 + j, f"={mo(MAT_START + i*5 + j, 'TIR')}", fmt=FMT_PCT2)
        ws.row_dimensions[r].height = 14   # matrices densas: 14 pt para que C + D + E compartan página
        r += 1
    cf_scale(ws, f"C{m0}:G{m0+4}")
    cf_active(ws, f"C{m0}:G{m0+4}", f"AND($B{m0}=0,C${h1}=0)", sides=("left", "right", "top", "bottom"))
    ws.row_dimensions[r].height = 6; r += 1
    hdr(ws, r, 2, 7, ["VAN proyecto @ tasa de descuento [USD]"] + [f"=INDEX(Mat_Tarifa,{j+1})" for j in range(5)], height=18)
    for j in range(5): ws.cell(row=r, column=3 + j).number_format = FMT_SIGNPCT
    h2 = r
    r += 1
    v0 = r
    for i in range(5):
        calc(ws, r, 2, f"=INDEX(Mat_CAPEX,{i+1})", fmt=FMT_SIGNPCT, bold=True, align="left")
        for j in range(5):
            calc(ws, r, 3 + j, f"={mo(MAT_START + i*5 + j, 'VAN')}", fmt=FMT_USD)
        ws.row_dimensions[r].height = 14   # matrices densas: 14 pt para que C + D + E compartan página
        r += 1
    cf_scale(ws, f"C{v0}:G{v0+4}", mid_num=0)
    cf_active(ws, f"C{v0}:G{v0+4}", f"AND($B{v0}=0,C${h2}=0)", sides=("left", "right", "top", "bottom"))
    note(ws, r, 2, '="El CAPEX fijo del Favorable ("&TEXT(INDEX(Esc_CAPEX_Fijo_Wp,4),"0.00")&" $/Wp, deck v4) equivale a "&TEXT(INDEX(Esc_CAPEX_Fijo_Wp,4)*Potencia_DC*1000/CAPEX_Base_f1-1,"+0%;-0%")&" de CAPEX sobre la base bottom-up; el Conservador, a +"&TEXT(INDEX(Esc_Factor_CAPEX,2)-1,"0%")&". Texto ladrillo: TIR bajo la tasa de descuento / VAN negativo."', c2=LC); ws.row_dimensions[r].height = 14   # matrices densas: 14 pt para que C + D + E compartan página
    r += 1
    ws.row_dimensions[r].height = 8; r += 1
    # ================================================================ D · Deuda
    section(ws, r, 2, LC, "D · Con deuda — TIR del accionista y DSCR mínimo: tasa (filas) × plazo (columnas); apalancamiento aparte", guide="borde = tasa y plazo (o apalancamiento) del Custom", guide_col=7)
    SENS_ANCHORS["D"] = f"B{r}"
    r += 1
    hdr(ws, r, 2, 5, ["TIR del accionista  ·  tasa ↓  plazo →"] + [f'=INDEX(Sens_Plazos,{j+1})&" años"' for j in range(3)], height=18)
    hdr(ws, r, 7, 10, ["DSCR mínimo"] + [f'=INDEX(Sens_Plazos,{j+1})&" años"' for j in range(3)], height=18)
    r += 1
    e0 = r
    for i in range(5):
        calc(ws, r, 2, f"=INDEX(Sens_Tasas,{i+1})", fmt=FMT_PCT2, bold=True, align="left"); calc(ws, r, 7, f"=INDEX(Sens_Tasas,{i+1})", fmt=FMT_PCT2, bold=True, align="left")
        for j in range(3):
            calc(ws, r, 3 + j, f"={mo(EQ_START + i*3 + j, 'TIR_eq')}", fmt=FMT_PCT2); calc(ws, r, 8 + j, f"={mo(EQ_START + i*3 + j, 'DSCR_min')}", fmt=FMT_X)
        ws.row_dimensions[r].height = 14   # matrices densas: 14 pt para que C + D + E compartan página
        r += 1
    cf_scale(ws, f"C{e0}:E{e0+4}")
    cf_text(ws, f"H{e0}:J{e0+4}", 1, "DSCR_Objetivo")
    cf_active(ws, f"C{e0}:E{e0+4}", f"AND($B{e0}=Tasa_Deuda,INDEX(Sens_Plazos,COLUMN()-2)=Plazo_Deuda)", sides=("left", "right", "top", "bottom"))
    cf_active(ws, f"H{e0}:J{e0+4}", f"AND($G{e0}=Tasa_Deuda,INDEX(Sens_Plazos,COLUMN()-7)=Plazo_Deuda)", sides=("left", "right", "top", "bottom"))
    ws.row_dimensions[r].height = 6; r += 1
    hdr(ws, r, 2, 6, ["Apalancamiento (deuda / CAPEX)", "TIR accionista", "DSCR mín", "Aporte de capital\n[USD]", "VAN accionista\n[USD]"], height=30)
    r += 1
    l0 = r
    for i in range(4):
        calc(ws, r, 2, f"=INDEX(Sens_Lev,{i+1})", fmt=FMT_PCT0, bold=True, align="left")
        calc(ws, r, 3, f"={mo(LEV_START+i,'TIR_eq')}", fmt=FMT_PCT2); calc(ws, r, 4, f"={mo(LEV_START+i,'DSCR_min')}", fmt=FMT_X)
        calc(ws, r, 5, f"={mo(LEV_START+i,'Aporte_eq')}", fmt=FMT_USD); calc(ws, r, 6, f"={mo(LEV_START+i,'VAN_eq')}", fmt=FMT_USD)
        ws.row_dimensions[r].height = 14   # matrices densas: 14 pt para que C + D + E compartan página
        r += 1
    cf_text(ws, f"D{l0}:D{l0+3}", 1, "DSCR_Objetivo")
    cf_active(ws, f"B{l0}:F{l0+3}", f"$B{l0}=Pct_Apalancamiento", sides=("top", "bottom"))
    callout(ws, r, 2, LC, '="Regla de bancabilidad típica: DSCR mínimo ≥ "&TEXT(DSCR_Objetivo,"0.00")&"x. Con "&Plazo_Deuda&" años y "&TEXT(Pct_Apalancamiento,"0%")&" el DSCR mínimo del Custom es "&TEXT(X_DSCR,"0.00")&"x en t = "&Anio_DSCR_Min&" (con tarifa plana y OPEX escalado el CFADS baja mientras la cuota es constante). Un plazo de 10 años lo mejora; 12 años no: las cuotas de los años 11-12 caen cuando termina la depreciación fiscal (10 años) → el plazo óptimo coincide con la vida fiscal de los equipos, o se sculpta la amortización. Menor apalancamiento también funciona (§H). El servicio de deuda no depende del escenario de energía; el DSCR sí (P90 lo reduce)."', height=44)
    r += 1
    ws.row_dimensions[r].height = 8; r += 1
    # ================================================================ E · Piso
    section(ws, r, 2, LC, "E · Escenario piso (todo conservador a la vez) vs Custom")
    SENS_ANCHORS["E"] = f"B{r}"
    r += 1
    hdr(ws, r, 2, 5, ["Indicador", "Custom", "Piso", "Δ (piso − Custom)"], height=18)
    caso_hdr(ws, r, 3, "X", size=SZ_TABLE)
    r += 1
    for lab, key, fmt in [("TIR proyecto", "TIR", FMT_PCT2), ("VAN proyecto", "VAN", FMT_USD), ("Payback (años)", "PB", FMT_YRS), ("LCOE ($/MWh)", "LCOE", FMT_DEC1), ("TIR accionista", "TIR_eq", FMT_PCT2), ("DSCR mínimo", "DSCR_min", FMT_X)]:
        label(ws, r, 2, lab); calc(ws, r, 3, f"={mo(CASE_BASE,key)}", fmt=fmt); calc(ws, r, 4, f"={mo(CASE_PISO,key)}", fmt=fmt); calc(ws, r, 5, f'=IFERROR(D{r}-C{r},"")', fmt=fmt, color=GRAFITO)
        ws.row_dimensions[r].height = 14   # matrices densas: 14 pt para que C + D + E compartan página
        r += 1
    callout(ws, r, 2, LC, '="Piso = P90 + CAPEX +"&TEXT(Sens_CAPEX,"0%")&" + OPEX +"&TEXT(Sens_OPEX_Up,"0%")&" + peaje "&TEXT(Sens_Peaje*100,"0.0")&" ¢/kWh + IVA no recuperable. Si el piso sigue siendo aceptable, el proyecto es robusto; si no, identifica qué palanca proteger contractualmente (precio EPC cerrado, O&M a precio fijo, cláusula de reapertura por peaje)."', height=30)
    r += 1
    ws.row_dimensions[r].height = 8; r += 1
    breaks.append(r - 1)
    # ================================================================ F · Dimensionamiento (transpuesto: métricas en filas, casos en columnas)
    section(ws, r, 2, LC, "F · Dimensionamiento — barridos por potencia DC y ratio DC/AC", guide="energía × recorte, CAPEX por drivers, terreno, deuda y Exergy recalculados en cada punto · ◆ y borde = Custom", guide_col=6)
    SENS_ANCHORS["F"] = f"B{r}"
    r += 1
    METRICS_P = [("Potencia DC", "kWp", FMT_INT, lambda ci: f"='{SM}'!${mcol(ci)}${PARAM_ROWS['P']}"),
                 ("Potencia AC", "kWac", FMT_INT, lambda ci: f"='{SM}'!${mcol(ci)}${SCAL_ROWS['AC']}"),
                 ("Terreno", "ha", "0.0", lambda ci: f"='{SM}'!${mcol(ci)}${SCAL_ROWS['ha']}"),
                 ("Energía año 1", "MWh", FMT_INT, lambda ci: f"={mo(ci,'E1')}"),
                 ("Cobertura del consumo", "%", FMT_PCT, lambda ci: f"={mo(ci,'Cob')}"),
                 ("Energía no reconocida Σ (art. 9)", "MWh", FMT_INT, lambda ci: f"={mo(ci,'NoRec')}"),
                 ("CAPEX industrial", "USD", FMT_USD, lambda ci: f"='{SM}'!${mcol(ci)}${SCAL_ROWS['K']}"),
                 ("CAPEX unitario", "$/Wp", FMT_WP, None),
                 ("TIR del proyecto", "%", FMT_PCT2, lambda ci: f"={mo(ci,'TIR')}"),
                 ("VAN del proyecto", "USD", FMT_USD, lambda ci: f"={mo(ci,'VAN')}"),
                 ("Payback desde COD", "años", FMT_YRS, lambda ci: f"={mo(ci,'PB')}"),
                 ("LCOE", "$/MWh", FMT_DEC1, lambda ci: f"={mo(ci,'LCOE')}"),
                 ("TIR del accionista", "%", FMT_PCT2, lambda ci: f"={mo(ci,'TIR_eq')}"),
                 ("DSCR mínimo", "x", FMT_X, lambda ci: f"={mo(ci,'DSCR_min')}"),
                 ("VAN negocio Exergy", "USD", FMT_USD, lambda ci: f"={mo(ci,'VAN_X')}"),
                 ("TIR consolidada del grupo", "%", FMT_PCT2, lambda ci: f"={mo(ci,'TIR_G')}")]

    def table_F(r, title, cases, head_f, metrics, unit_wp_from=("CAPEX industrial", "Potencia DC"), anchor_key=None, active_formula=None):
        """Tabla transpuesta: cabecera con un caso por columna (C…), una métrica por fila. Devuelve (fila de cabecera, dict métrica→fila, última fila)."""
        n = len(cases)
        hdr_cases(ws, r, 3, [head_f(i) for i in range(n)], height=30)
        c = ws.cell(row=r, column=2, value=title); c.font = Font(name=FONT, size=SZ_TABLE, bold=True, color=GRAFITO); c.alignment = Alignment(vertical="center", wrap_text=True)
        hr = r
        r += 1
        rows_of = {}
        for lab, u, fmt, fn in metrics:
            label(ws, r, 2, lab, bold=(lab in ("TIR del proyecto", "VAN del proyecto")))
            for i, ci in enumerate(cases):
                cc = 3 + i
                if fn is None:   # CAPEX unitario = CAPEX / kWp (misma fórmula que la v1.2)
                    f = f"={col(cc)}{rows_of[unit_wp_from[0]]}/({col(cc)}{rows_of[unit_wp_from[1]]}*1000)"
                elif callable(fn):
                    f = fn(ci) if fn.__code__.co_argcount == 1 else fn(ci, cc, rows_of)
                calc(ws, r, cc, f, fmt=fmt, bold=(lab in ("TIR del proyecto", "VAN del proyecto")))
            c = ws.cell(row=r, column=3 + n, value=u); c.font = font(size=SZ_TABLE, color=GRAFITO); c.alignment = Alignment(horizontal="left", vertical="center"); c.border = B_BOTTOM
            ws.row_dimensions[r].height = RH
            rows_of[lab] = r
            r += 1
        last_col = 3 + n
        if active_formula:
            cf_active(ws, f"C{hr}:{col(2+n)}{r-1}", active_formula)
        if anchor_key:
            SENS_ANCHORS[anchor_key] = f"B{hr}"
        return hr, rows_of, r - 1

    # F.1 potencia
    SENS_ANCHORS["F1"] = f"B{r}"
    casesP = [SWEEP_P_START + i for i in range(N_SWEEP_P)]
    hF1, rowsF1, lastF1 = table_F(r, "F.1 · Potencia DC (ratio del Custom) →", casesP,
                                  lambda i: f'=IF(INDEX(Sweep_P,{i+1})=Potencia_DC,"◆ ","")&TEXT(INDEX(Sweep_P,{i+1})/1000,"0.0")&" MWp"', METRICS_P,
                                  active_formula=f"INDEX(Sweep_P,COLUMN()-2)=Potencia_DC")
    r = lastF1 + 1
    fmtP = lambda key: [m for m in METRICS_P if m[0] == key][0]
    cf_scale(ws, f"C{rowsF1['TIR del proyecto']}:{col(2+N_SWEEP_P)}{rowsF1['TIR del proyecto']}")
    cf_scale(ws, f"C{rowsF1['VAN del proyecto']}:{col(2+N_SWEEP_P)}{rowsF1['VAN del proyecto']}", mid_num=0)
    ws.conditional_formatting.add(f"C{rowsF1['Energía no reconocida Σ (art. 9)']}:{col(2+N_SWEEP_P)}{rowsF1['Energía no reconocida Σ (art. 9)']}", CellIsRule(operator="greaterThan", formula=["0"], font=Font(color=LADRILLO, bold=True)))
    note(ws, r, 2, '="Candados: art. 9 (energía no reconocida > 0 → la demanda de "&TEXT(Consumo_Anual/1000,"#,##0")&" MWh/año se alcanza en ≈ "&TEXT(Consumo_Anual/Yield_Ref/F_Recorte,"#,##0")&" kWp), excedentes mensuales desde ≈ 7,4 MWp (junio, 04), alimentador ("&TEXT(Capacidad_Alimentador_kW,"#,##0")&" kW → ≈ "&TEXT(Capacidad_Alimentador_kW*Ratio_DCAC,"#,##0")&" kWp) y predio ("&TEXT(Ha_Disponibles*Densidad_MWp_ha*1000,"#,##0")&" kWp). La potencia de 01 es una lista 5–8 MWp; CAPEX por drivers: pesos «por confirmar», validez 3–8 MWp."', c2=LC)
    fit_row(ws, r, [("x" * 320, 150, 8.5)])
    r += 1
    ws.row_dimensions[r].height = 8; r += 1
    # gráficos F.1 (bloque auxiliar en U:W, fuera del área de impresión)
    ax0 = hF1 + 1
    ws.column_dimensions["U"].width = 30; ws.column_dimensions["V"].width = 12; ws.column_dimensions["W"].width = 12
    c = ws.cell(row=ax0 - 1, column=21, value="Auxiliares de los gráficos §F (no editar) · ejes fijados a 4,5–8,5 MWp y 1,05–1,55: si edita los barridos fuera de ese rango, amplíe el eje"); c.font = font(size=8.5, color=GRAFITO)
    kwp1, van1, tir1, tireq1, tirg1, vanx1 = (rowsF1[k] for k in ("Potencia DC", "VAN del proyecto", "TIR del proyecto", "TIR del accionista", "TIR consolidada del grupo", "VAN negocio Exergy"))
    aux = [("Custom F.1 · TIR (x en MWp)", "=Potencia_DC/1000", f"={mo(CASE_BASE,'TIR')}", "0.00", FMT_PCT2),
           ("Custom F.1 · VAN (x en MWp)", "=Potencia_DC/1000", f"={mo(CASE_BASE,'VAN')}", "0.00", FMT_USD),
           ("Referencia tasa de descuento · x1 (MWp)", "=INDEX(Sweep_P,1)/1000", "=Tasa_Descuento", "0.00", FMT_PCT2),
           ("Referencia tasa de descuento · x2 (MWp)", f"=INDEX(Sweep_P,{N_SWEEP_P})/1000", "=Tasa_Descuento", "0.00", FMT_PCT2),
           ("Límite art. 9 · y1 (x en MWp)", "=Consumo_Anual/Yield_Ref/F_Recorte/1000", f"=MIN(C{van1}:{col(2+N_SWEEP_P)}{van1})", "0.00", FMT_USD),
           ("Límite art. 9 · y2 (x en MWp)", "=Consumo_Anual/Yield_Ref/F_Recorte/1000", f"=MAX(C{van1}:{col(2+N_SWEEP_P)}{van1})", "0.00", FMT_USD),
           ("Custom F.2 · TIR", "=Ratio_DCAC", f"={mo(CASE_BASE,'TIR')}", "0.00", FMT_PCT2),
           ("Custom F.2 · LCOE", "=Ratio_DCAC", f"={mo(CASE_BASE,'LCOE')}", "0.00", FMT_DEC1),
           ("Referencia tarifa evitable · x1", "=INDEX(Sweep_Ratio,1)", "=Tarifa_MWh", "0.00", FMT_DEC1),
           ("Referencia tarifa evitable · x2", f"=INDEX(Sweep_Ratio,{N_SWEEP_R})", "=Tarifa_MWh", "0.00", FMT_DEC1)]
    for i in range(N_SWEEP_R):
        aux.append((f"Ratio paso {i+1} (eje x de F.2)", f"=INDEX(Sweep_Ratio,{i+1})", None, "0.00", None))
    for i in range(N_SWEEP_P):
        aux.append((f"Potencia paso {i+1} en MWp (eje x de F.1)", f"={col(3+i)}{kwp1}/1000", None, "0.00", None))
    for k, (lab, fxv, fyv, fmx, fmy) in enumerate(aux):
        rr = ax0 + k
        c = ws.cell(row=rr, column=21, value=lab); c.font = font(size=8.5, color=GRAFITO)
        c = ws.cell(row=rr, column=22, value=fxv); c.font = font(size=8.5, color=GRAFITO); c.number_format = fmx
        if fyv is not None:
            c = ws.cell(row=rr, column=23, value=fyv); c.font = font(size=8.5, color=GRAFITO); c.number_format = fmy
    AX = lambda r0, r1=None: Reference(ws, min_col=22, min_row=r0, max_row=(r1 or r0))
    AY = lambda r0, r1=None: Reference(ws, min_col=23, min_row=r0, max_row=(r1 or r0))
    xP = AX(ax0 + 10 + N_SWEEP_R, ax0 + 10 + N_SWEEP_R + N_SWEEP_P - 1)
    ratio_x = AX(ax0 + 10, ax0 + 10 + N_SWEEP_R - 1)
    ROWP = lambda row: Reference(ws, min_col=3, max_col=2 + N_SWEEP_P, min_row=row)
    ROWR = lambda row: Reference(ws, min_col=3, max_col=2 + N_SWEEP_R, min_row=row)
    CW, CH, NCHF = 11.8, 9.0, 19     # dos gráficos por fila: B (termina en F) y G (termina antes de L: 613 + 446 px < 1.105 px); 19 × (15 − 1) = 266 pt ≥ 9,0 cm (255 pt) + 6
    gr = r
    chart_scatter(ws, f"B{gr}", "F.1 · VAN del proyecto y de Exergy vs potencia DC", [
        dict(x=xP, y=ROWP(van1), name="VAN proyecto", color=TERRACOTA, width=2.25, labels=[N_SWEEP_P - 1], label_fmt='#,##0', label_pos="t"),
        dict(x=xP, y=ROWP(vanx1), name="VAN Exergy", color=PIEDRA, width=1.5, labels=[N_SWEEP_P - 1], label_fmt='#,##0', label_pos="b"),
        dict(x=AX(ax0 + 4, ax0 + 5), y=AY(ax0 + 4, ax0 + 5), name="límite art. 9", kind="ref", color=ARCILLA),
        dict(x=AX(ax0 + 1), y=AY(ax0 + 1), name="Custom", kind="point", color=X_COL, labels=[0], label_fmt='#,##0', label_pos="r")],
        w=CW, h=CH, x_fmt='0.0', y_fmt=FMT_K, x_title="potencia DC (MWp)", legend="b", x_min=4.5, x_max=8.5)
    chart_scatter(ws, f"G{gr}", "F.1 · TIR vs potencia DC", [
        dict(x=xP, y=ROWP(tir1), name="TIR proyecto", color=TERRACOTA, width=2.25, labels=[N_SWEEP_P - 1], label_fmt='0.0%', label_pos="t"),
        dict(x=xP, y=ROWP(tireq1), name="TIR accionista", color=GRAFITO, width=1.5, labels=[N_SWEEP_P - 1], label_fmt='0.0%', label_pos="t"),
        dict(x=xP, y=ROWP(tirg1), name="TIR grupo", color=PIEDRA, width=1.25, labels=[N_SWEEP_P - 1], label_fmt='0.0%', label_pos="b"),
        dict(x=AX(ax0 + 2, ax0 + 3), y=AY(ax0 + 2, ax0 + 3), name="tasa de descuento", kind="ref"),
        dict(x=AX(ax0), y=AY(ax0), name="Custom", kind="point", color=X_COL, labels=[0], label_fmt='0.0%', label_pos="b")],
        w=CW, h=CH, x_fmt='0.0', y_fmt='0%', x_title="potencia DC (MWp)", legend="b", x_min=4.5, x_max=8.5, y_min=0.05)
    for k in range(gr, gr + NCHF):
        ws.row_dimensions[k].height = 15
    r = gr + NCHF
    breaks.append(r - 1)
    # F.2 ratio (DC activa)
    METRICS_R = [("Potencia DC", "kWp", FMT_INT, lambda ci: f"='{SM}'!${mcol(ci)}${PARAM_ROWS['P']}"),
                 ("Potencia AC", "kWac", FMT_INT, lambda ci: f"='{SM}'!${mcol(ci)}${SCAL_ROWS['AC']}"),
                 ("Factor de recorte", "×", "0.0000", lambda ci: f"='{SM}'!${mcol(ci)}${SCAL_ROWS['frec']}"),
                 ("Energía año 1", "MWh", FMT_INT, lambda ci: f"={mo(ci,'E1')}"),
                 ("Δ energía vs Custom", "%", FMT_SIGNPCT, lambda ci, cc, rows_of: f"={col(cc)}{rows_of['Energía año 1']}/{mo(CASE_BASE,'E1')}-1"),
                 ("CAPEX industrial", "USD", FMT_USD, lambda ci: f"='{SM}'!${mcol(ci)}${SCAL_ROWS['K']}"),
                 ("Δ CAPEX vs Custom", "%", FMT_SIGNPCT, lambda ci, cc, rows_of: f"={col(cc)}{rows_of['CAPEX industrial']}/'{SM}'!$B${SCAL_ROWS['K']}-1"),
                 ("CAPEX unitario", "$/Wp", FMT_WP, None),
                 ("TIR del proyecto", "%", FMT_PCT2, lambda ci: f"={mo(ci,'TIR')}"),
                 ("VAN del proyecto", "USD", FMT_USD, lambda ci: f"={mo(ci,'VAN')}"),
                 ("Payback desde COD", "años", FMT_YRS, lambda ci: f"={mo(ci,'PB')}"),
                 ("LCOE", "$/MWh", FMT_DEC1, lambda ci: f"={mo(ci,'LCOE')}"),
                 ("TIR del accionista", "%", FMT_PCT2, lambda ci: f"={mo(ci,'TIR_eq')}"),
                 ("DSCR mínimo", "x", FMT_X, lambda ci: f"={mo(ci,'DSCR_min')}"),
                 ("VAN negocio Exergy", "USD", FMT_USD, lambda ci: f"={mo(ci,'VAN_X')}"),
                 ("Potencia AC vs alimentador", "", None, lambda ci, cc, rows_of: f'=IF({col(cc)}{rows_of["Potencia AC"]}<=Capacidad_Alimentador_kW,"● cabe","■ excede")')]
    casesR = [SWEEP_R_START + i for i in range(N_SWEEP_R)]
    hF2, rowsF2, lastF2 = table_F(r, "F.2 · Ratio DC/AC (potencia DC del Custom) →", casesR,
                                  lambda i: f'=IF(ABS(INDEX(Sweep_Ratio,{i+1})-Ratio_DCAC)<0.001,"◆ ","")&"ratio "&TEXT(INDEX(Sweep_Ratio,{i+1}),"0.00")', METRICS_R,
                                  anchor_key="F2", active_formula=f"ABS(INDEX(Sweep_Ratio,COLUMN()-2)-Ratio_DCAC)<0.001")
    r = lastF2 + 1
    for cc in range(3, 3 + N_SWEEP_R):
        cell = ws.cell(row=rowsF2["Potencia AC vs alimentador"], column=cc); cell.alignment = Alignment(horizontal="right", vertical="center"); cell.font = font(size=8.5, color=GRAFITO)
    cf_scale(ws, f"C{rowsF2['TIR del proyecto']}:{col(2+N_SWEEP_R)}{rowsF2['TIR del proyecto']}")
    cf_scale(ws, f"C{rowsF2['TIR del accionista']}:{col(2+N_SWEEP_R)}{rowsF2['TIR del accionista']}")
    _glyph_rng = f"C{rowsF2['Potencia AC vs alimentador']}:{col(2+N_SWEEP_R)}{rowsF2['Potencia AC vs alimentador']}"
    ws.conditional_formatting.add(_glyph_rng, FormulaRule(formula=[f'LEFT(C{rowsF2["Potencia AC vs alimentador"]},1)="■"'], font=Font(color=LADRILLO, bold=True)))
    ws.conditional_formatting.add(_glyph_rng, FormulaRule(formula=[f'LEFT(C{rowsF2["Potencia AC vs alimentador"]},1)="●"'], font=Font(color=SALVIA, bold=True)))
    note(ws, r, 2, "Subir el ratio abarata inversores y transformación (drivers Wac) y reduce la potencia AC (más holgura en el alimentador) a cambio de recorte de energía (curva pvlib en 04). Con el recurso de Montecristi (Pdc máx ≈ 0,90 kW/kWp) el recorte es pequeño hasta 1,5.", c2=LC)
    fit_row(ws, r, [("x" * 250, 150, 8.5)])
    r += 1
    ws.row_dimensions[r].height = 8; r += 1
    tir2, lcoe2 = rowsF2["TIR del proyecto"], rowsF2["LCOE"]
    gr = r
    chart_scatter(ws, f"B{gr}", "F.2 · TIR del proyecto vs ratio DC/AC", [
        dict(x=ratio_x, y=ROWR(tir2), name="TIR proyecto", color=TERRACOTA, width=2.25, labels=[0, N_SWEEP_R - 1], label_fmt='0.00%', label_pos="t"),
        dict(x=AX(ax0 + 6), y=AY(ax0 + 6), name="Custom", kind="point", color=X_COL, labels=[0], label_fmt='0.00%', label_pos="b")],
        w=CW, h=CH, x_fmt='0.00', y_fmt='0.0%', x_title="ratio DC/AC", legend="b", x_min=1.05, x_max=1.55)
    chart_scatter(ws, f"G{gr}", "F.2 · LCOE vs ratio DC/AC", [
        dict(x=ratio_x, y=ROWR(lcoe2), name="LCOE", color=TERRACOTA, width=2.25, labels=[0, N_SWEEP_R - 1], label_fmt='0.0', label_pos="t"),
        dict(x=AX(ax0 + 8, ax0 + 9), y=AY(ax0 + 8, ax0 + 9), name="tarifa evitable", kind="ref", labels=[1], label_fmt='0.0', label_pos="t"),
        dict(x=AX(ax0 + 7), y=AY(ax0 + 7), name="Custom", kind="point", color=X_COL, labels=[0], label_fmt='0.0', label_pos="b")],
        w=CW, h=CH, x_fmt='0.00', y_fmt='0.0', x_title="ratio DC/AC", legend="b", x_min=1.05, x_max=1.55)
    for k in range(gr, gr + NCHF):
        ws.row_dimensions[k].height = 15
    r = gr + NCHF
    breaks.append(r - 1)
    # F.3 AC fija
    casesA = [SWEEP_A_START + i for i in range(N_SWEEP_R)]
    hF3, rowsF3, lastF3 = table_F(r, '="F.3 · AC fija = "&TEXT(Sweep_AC_Fija,"#,##0")&" kWac; DC variable →"', casesA,
                                  lambda i: f'="ratio "&TEXT(INDEX(Sweep_Ratio,{i+1}),"0.00")', METRICS_P, anchor_key="F3")
    r = lastF3 + 1
    cf_scale(ws, f"C{rowsF3['TIR del proyecto']}:{col(2+N_SWEEP_R)}{rowsF3['TIR del proyecto']}")
    cf_scale(ws, f"C{rowsF3['VAN del proyecto']}:{col(2+N_SWEEP_R)}{rowsF3['VAN del proyecto']}", mid_num=0)
    note(ws, r, 2, "Responde a «¿cuánto DC pongo detrás del mismo punto de conexión?»: la potencia AC se fija y el DC crece con el ratio.", c2=LC); ws.row_dimensions[r].height = RH
    r += 1
    ws.row_dimensions[r].height = 8; r += 1
    # ================================================================ H · Deuda máxima (comparte página con F.3)
    section(ws, r, 2, LC, "H · Deuda máxima sostenible — apalancamiento que cumple el DSCR objetivo, por plazo", guide="interpolación lineal sobre la malla Sweep_Lev; texto ladrillo = DSCR bajo 1,00x, arcilla = bajo el objetivo", guide_col=6)
    SENS_ANCHORS["H"] = f"B{r}"
    r += 1
    hdr(ws, r, 2, 2 + N_SWEEP_LEV, ["DSCR mínimo  ·  plazo ↓  apalancamiento →"] + [f"=INDEX(Sweep_Lev,{i+1})" for i in range(N_SWEEP_LEV)], height=30)
    for i in range(N_SWEEP_LEV): ws.cell(row=r, column=3 + i).number_format = FMT_PCT0
    hdr(ws, r, 3 + N_SWEEP_LEV, LC, ["Deuda máx.\n(% CAPEX)", "Deuda máx.\n[USD]", "Aporte mín.\n[USD]", "TIR accionista\na esa deuda"], height=30)
    hH = r
    r += 1
    h0 = r
    for j in range(3):
        calc(ws, r, 2, f'="Plazo "&INDEX(Sens_Plazos,{j+1})&" años"', bold=True, align="left")
        for i in range(N_SWEEP_LEV):
            calc(ws, r, 3 + i, f"={mo(DM_START + j*N_SWEEP_LEV + i, 'DSCR_min')}", fmt=FMT_X)
        rngD = f"$C{r}:${col(2+N_SWEEP_LEV)}{r}"
        k_ = f"COUNTIF({rngD},\">=\"&DSCR_Objetivo)"
        interp = (f"IF({k_}=0,0,IF({k_}>={N_SWEEP_LEV},INDEX(Sweep_Lev,{N_SWEEP_LEV}),"
                  f"INDEX(Sweep_Lev,{k_})+(INDEX(Sweep_Lev,{k_}+1)-INDEX(Sweep_Lev,{k_}))*(INDEX({rngD},{k_})-DSCR_Objetivo)/(INDEX({rngD},{k_})-INDEX({rngD},{k_}+1))))")
        cD = 3 + N_SWEEP_LEV     # H
        calc(ws, r, cD, "=" + interp, fmt=FMT_PCT, bold=True)
        calc(ws, r, cD + 1, f"={col(cD)}{r}*CAPEX_Total", fmt=FMT_USD); calc(ws, r, cD + 2, f"=CAPEX_Total+Terreno_SALELGI+IVA_Total*Fase_m1-{col(cD+1)}{r}", fmt=FMT_USD)
        tir_rng = f"'{SM}'!${mcol(DM_START + j*N_SWEEP_LEV)}${OUT_ROWS['TIR_eq']}:${mcol(DM_START + j*N_SWEEP_LEV + N_SWEEP_LEV - 1)}${OUT_ROWS['TIR_eq']}"
        calc(ws, r, cD + 3, f'=IFERROR(IF({k_}=0,"n/a",IF({k_}>={N_SWEEP_LEV},INDEX({tir_rng},{N_SWEEP_LEV}),INDEX({tir_rng},{k_})+(INDEX({tir_rng},{k_}+1)-INDEX({tir_rng},{k_}))*({col(cD)}{r}-INDEX(Sweep_Lev,{k_}))/(INDEX(Sweep_Lev,{k_}+1)-INDEX(Sweep_Lev,{k_})))),"n/a")', fmt=FMT_PCT2)
        ws.row_dimensions[r].height = RH
        r += 1
    cf_text(ws, f"C{h0}:{col(2+N_SWEEP_LEV)}{h0+2}", 1, "DSCR_Objetivo")
    cf_active(ws, f"B{h0}:{col(LC)}{h0+2}", f"INDEX(Sens_Plazos,ROW()-{h0}+1)=Plazo_Deuda", sides=("top", "bottom"))
    name(wb, "Deuda_Max_Plazo1", S10, f"${col(cD)}${h0}"); name(wb, "Deuda_Max_Plazo2", S10, f"${col(cD)}${h0+1}"); name(wb, "Deuda_Max_Plazo3", S10, f"${col(cD)}${h0+2}")
    # fila de referencia para el gráfico (objetivo constante)
    label(ws, r, 2, "DSCR objetivo (referencia del gráfico)", size=9, color=GRAFITO)   # label() acepta color
    for i in range(N_SWEEP_LEV):
        calc(ws, r, 3 + i, "=DSCR_Objetivo", fmt=FMT_X, color=GRAFITO, size=9)
    ws.row_dimensions[r].height = RH
    ref_r = r
    r += 1
    note(ws, r, 2, '="Deuda máxima = mayor apalancamiento cuyo DSCR mínimo ≥ "&TEXT(DSCR_Objetivo,"0.00")&"x (a la tasa activa "&TEXT(Tasa_Deuda,"0.0%")&"), interpolado entre los puntos de la malla; 0 % = ni el menor apalancamiento cumple. Aporte mínimo = CAPEX + terreno (si SALELGI) + IVA del tramo −1 − deuda."', c2=LC)
    fit_row(ws, r, [("x" * 260, 150, 8.5)])
    r += 1
    ws.row_dimensions[r].height = 8; r += 1
    gr = r
    catsH = Reference(ws, min_col=3, max_col=2 + N_SWEEP_LEV, min_row=hH)
    chart_lines(ws, f"B{gr}", "DSCR mínimo por apalancamiento y plazo · la deuda máxima es el cruce con el objetivo", catsH,
                [dict(ref=Reference(ws, min_col=3, max_col=2 + N_SWEEP_LEV, min_row=h0 + j), name="", color=colr, width=w_, marker="circle", marker_size=5)
                 for j, (colr, w_) in enumerate(((PIEDRA, 1.5), (TERRACOTA, 2.25), (GRAFITO, 1.5)))],
                ref=dict(ref=Reference(ws, min_col=3, max_col=2 + N_SWEEP_LEV, min_row=ref_r), name="DSCR objetivo", color=ARCILLA, width=1.5, dash="dash"),
                w=19.0, h=6.0, y_fmt='0.00"x"', legend="r", from_rows=True, x_title="apalancamiento (deuda / CAPEX)")
    # nombres de serie vivos: "plazo N años" desde Sens_Plazos
    ch = ws._charts[-1]
    from openpyxl.chart.series import SeriesLabel
    from openpyxl.chart.data_source import StrRef
    for j in range(3):
        ch.series[j].tx = SeriesLabel(strRef=StrRef(f"'{S10}'!$B${h0 + j}"))
    NCHH = 13
    for k in range(gr, gr + NCHH):
        ws.row_dimensions[k].height = 15          # 13 × 14 = 182 ≥ 6,0 cm (170 pt)
    r = gr + NCHH
    breaks.append(r - 1)
    # ================================================================ G · Terreno
    section(ws, r, 2, LC, "G · Quién compra el terreno — Exergy (y arrienda) vs SALELGI (compra)", guide="ambas alternativas con los demás supuestos del Custom; el comprador se elige en 01", guide_col=6)
    SENS_ANCHORS["G"] = f"B{r}"
    r += 1
    hdr(ws, r, 2, LC, ["Indicador", '=IF(Comprador_Terreno="Exergy","◆ Exergy compra","Exergy compra")', '=IF(Comprador_Terreno="SALELGI","◆ SALELGI compra","SALELGI compra")', "Δ (S − E)", "Lectura"] + [""] * 6, height=30)
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=LC)
    ws.cell(row=r, column=6).alignment = Alignment(horizontal="left", vertical="center")
    hG = r
    r += 1

    def mo_sel(sel_expr, key):
        return f"IF('{SM}'!$B${PARAM_ROWS['terr']}=0,{mo(CASE_BASE,key) if sel_expr=='E' else mo(CASE_TERR_ALT,key)},{mo(CASE_TERR_ALT,key) if sel_expr=='E' else mo(CASE_BASE,key)})"
    g_rows = [("CAPEX total SALELGI incl. terreno [USD]", "K_T", FMT_USD, "El terreno entra al CAPEX de SALELGI con costos de transacción; no se deprecia ni paga fee."),
              ("Aporte de capital con deuda [USD]", "Aporte_eq", FMT_USD, "Sube ≈ el costo del terreno si el banco no lo financia (Deuda_Financia_Terreno)."),
              ("OPEX año 1 SALELGI [USD]", "OPEX1", FMT_USD, "Sin arriendo; con predial."),
              ("TIR del proyecto (sin deuda)", "TIR", FMT_PCT2, "SALELGI paga hoy el terreno pero se ahorra el arriendo y lo recupera al año 25."),
              ("VAN del proyecto @ tasa de descuento [USD]", "VAN", FMT_USD, ""),
              ("TIR del accionista (con deuda)", "TIR_eq", FMT_PCT2, ""), ("DSCR mínimo", "DSCR_min", FMT_X, "Mejora sin arriendo (CFADS mayor) si la deuda no crece."),
              ("VAN negocio Exergy [USD]", "VAN_X", FMT_USD, "Exergy pierde la línea terreno (compra, arriendo, predial, residual)."),
              ("Ingreso neto nominal Exergy Σ [USD]", "Nominal_X", FMT_USD, ""),
              ("TIR consolidada del grupo", "TIR_G", FMT_PCT2, "El grupo sólo cambia por impuestos y timing: el terreno es el mismo activo en manos distintas.")]
    g0 = r
    for lab, key, fmt, lect in g_rows:
        label(ws, r, 2, lab, bold=(key in ("TIR", "VAN_X", "TIR_G")))
        if key == "K_T":
            fE = f"IF('{SM}'!$B${PARAM_ROWS['terr']}=0,'{SM}'!$B${SCAL_ROWS['K']},'{SM}'!${mcol(CASE_TERR_ALT)}${SCAL_ROWS['K']})"
            fS = f"IF('{SM}'!$B${PARAM_ROWS['terr']}=1,'{SM}'!$B${SCAL_ROWS['K']}+'{SM}'!$B${SCAL_ROWS['Terr']},'{SM}'!${mcol(CASE_TERR_ALT)}${SCAL_ROWS['K']}+'{SM}'!${mcol(CASE_TERR_ALT)}${SCAL_ROWS['Terr']})"
        elif key == "OPEX1":
            fE = f"IF('{SM}'!$B${PARAM_ROWS['terr']}=0,'{SM}'!$B${SCAL_ROWS['OPEX1']},'{SM}'!${mcol(CASE_TERR_ALT)}${SCAL_ROWS['OPEX1']})"
            fS = f"IF('{SM}'!$B${PARAM_ROWS['terr']}=1,'{SM}'!$B${SCAL_ROWS['OPEX1']},'{SM}'!${mcol(CASE_TERR_ALT)}${SCAL_ROWS['OPEX1']})"
        else:
            fE = mo_sel("E", key); fS = mo_sel("S", key)
        calc(ws, r, 3, "=" + fE, fmt=fmt); calc(ws, r, 4, "=" + fS, fmt=fmt); calc(ws, r, 5, f'=IFERROR(D{r}-C{r},"")', fmt=fmt, color=GRAFITO)
        note(ws, r, 6, lect, border=True, valign="center", c2=LC)
        for cc in range(6, LC + 1): ws.cell(row=r, column=cc).border = B_BOTTOM
        ws.row_dimensions[r].height = RH
        r += 1
    cf_active(ws, f"C{g0}:C{r-1}", 'Comprador_Terreno="Exergy"'); cf_active(ws, f"D{g0}:D{r-1}", 'Comprador_Terreno="SALELGI"')
    ws.row_dimensions[r].height = 6; r += 1
    hdr(ws, r, 2, 8, ["Precio del terreno ($/ha) — sensibilidad", "", "SALELGI compra: TIR", "SALELGI compra: VAN acc.", "Exergy compra: TIR", "Exergy compra: VAN Exergy", "TIR grupo (S / E)"], height=30)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    r += 1
    p0 = r
    for i, fp in enumerate(PRECIO_STEPS):
        calc(ws, r, 2, f"=Precio_Terreno_ha*{fp}", fmt=FMT_USD, bold=True, align="left"); calc(ws, r, 3, f'="× {fp}"', align="center", color=GRAFITO)
        calc(ws, r, 4, f"={mo(TS_START+i,'TIR')}", fmt=FMT_PCT2); calc(ws, r, 5, f"={mo(TS_START+i,'VAN_eq')}", fmt=FMT_USD)
        calc(ws, r, 6, f"={mo(TX_START+i,'TIR')}", fmt=FMT_PCT2); calc(ws, r, 7, f"={mo(TX_START+i,'VAN_X')}", fmt=FMT_USD)
        calc(ws, r, 8, f'=TEXT({mo(TS_START+i,"TIR_G")},"0.00%")&" / "&TEXT({mo(TX_START+i,"TIR_G")},"0.00%")', align="center")
        ws.row_dimensions[r].height = RH
        r += 1
    cf_active(ws, f"B{p0}:H{p0+len(PRECIO_STEPS)-1}", f"$B{p0}=Precio_Terreno_ha", sides=("top", "bottom"))
    callout(ws, r, 2, LC, '="Lectura: con renta "&TEXT(Renta_Terreno_ha,"#,##0")&" $/ha ("&IFERROR(TEXT(Renta_Terreno_ha/Precio_Terreno_ha,"0%"),"n/a")&" del precio) y tasa de descuento "&TEXT(Tasa_Descuento,"0%")&", la línea terreno es casi neutra: para SALELGI comprar equivale a prepagar 25 años de arriendo con recuperación del capital; la diferencia la hacen los costos de transacción, el tratamiento fiscal del arriendo (deducible) y quién financia. Si la renta supera la tasa de descuento, a SALELGI le conviene comprar; si es menor, arrendar."', height=44)
    r += 1
    ws.row_dimensions[r].height = 8; r += 1
    # ---- v3.0 (V6): quién paga el reemplazo de inversores (capa de diseño; caso «Reemplazo alt.» del Motor)
    section(ws, r, 2, LC, "G.2 · Quién paga el reemplazo de inversores — SALELGI (capex depreciable) vs Exergy (reserva del fee de O&M)", guide="el pagador se elige en 01 (Reemplazo_Pagador); el alterno se calcula en el Motor", guide_col=7)
    r += 1
    hdr(ws, r, 2, LC, ["Indicador", '=IF(Reemplazo_Pagador="SALELGI","◆ SALELGI paga","SALELGI paga")', '=IF(Reemplazo_Pagador="Exergy","◆ Exergy paga","Exergy paga")', "Δ (E − S)", "Lectura"] + [""] * 6, height=30)
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=LC)
    ws.cell(row=r, column=6).alignment = Alignment(horizontal="left", vertical="center")
    r += 1
    rep_par = f"'{SM}'!$B${PARAM_ROWS['rep']}"
    def mo_rep(who, key):
        # SALELGI paga: Custom si rep = 1, el alterno si rep ≠ 1 (rep = 0 → alterno = SALELGI) · Exergy paga: Custom si rep = 2, el alterno si rep = 1, n/a si rep = 0
        if who == "S":
            return f"IF({rep_par}=1,{mo(CASE_BASE,key)},{mo(T_REP,key)})"
        return f'IF({rep_par}=2,{mo(CASE_BASE,key)},IF({rep_par}=1,{mo(T_REP,key)},"n/a"))'
    rep_rows = [("TIR del proyecto (sin deuda)", "TIR", FMT_PCT2, '="El reemplazo ("&TEXT(Reemplazo_USD,"#,##0")&" USD en t = "&Reemplazo_Anio&") es capex de SALELGI si lo paga ella; Exergy no altera el flujo de SALELGI."'),
                ("VAN del proyecto @ tasa de descuento [USD]", "VAN", FMT_USD, ""),
                ("TIR del accionista (con deuda)", "TIR_eq", FMT_PCT2, ""),
                ("Payback desde COD [años]", "PB", FMT_YRS, "Con el reemplazo el acumulado puede volver a caer: el payback toma el último cruce."),
                ("VAN negocio Exergy [USD]", "VAN_X", FMT_USD, "Si paga Exergy, el gasto sale de la reserva incluida en el fee de O&M (deducible en Exergy)."),
                ("TIR consolidada del grupo", "TIR_G", FMT_PCT2, "El grupo paga el reemplazo en cualquier caso: sólo cambian los impuestos y quién lo asume.")]
    g2 = r
    for lab, key, fmt, lect in rep_rows:
        label(ws, r, 2, lab, bold=(key in ("TIR", "VAN_X")))
        calc(ws, r, 3, "=" + mo_rep("S", key), fmt=fmt); calc(ws, r, 4, "=" + mo_rep("E", key), fmt=fmt); calc(ws, r, 5, f'=IFERROR(D{r}-C{r},"")', fmt=fmt, color=GRAFITO)
        note(ws, r, 6, lect, border=True, valign="center", c2=LC)
        for cc in range(6, LC + 1): ws.cell(row=r, column=cc).border = B_BOTTOM
        ws.row_dimensions[r].height = RH
        r += 1
    cf_active(ws, f"C{g2}:C{r-1}", 'Reemplazo_Pagador="SALELGI"'); cf_active(ws, f"D{g2}:D{r-1}", 'Reemplazo_Pagador="Exergy"')
    ws.row_dimensions[r].height = 8; r += 1
    # ---- impresión
    for b in breaks:
        ws.row_breaks.append(Break(id=b))
    setup_print(ws, landscape=True, scale=82, title_rows=None, area=f"A1:{col(LC)}{r}")
    return ws, (t0, t1)

# -*- coding: utf-8 -*-
"""Orquestador v3.0 (ronda 2 de motor sobre la arquitectura de casos v2.0). ESC_DEF=v20 construye el bloque B con la definición v2.0 y los
parámetros de la ronda 2 en neutro (prueba de invariancia R1 frente a gen_v20_calc); v13 = definición v1.3 (prueba histórica)."""
import sys
from openpyxl import Workbook
import xl_helpers as XH
from xl_helpers import *
from build_core import *
from build_content import *
from build_sens30 import build_sensibilidad
from build_guia import build_guia, SG
from build_supuestos30 import build_supuestos
from build_portada30 import build_portada

OUT = sys.argv[1] if len(sys.argv) > 1 else "/root/gpm13/out20/Modelo_FV_5MWp_GPM_v2.0.xlsx"


def main():
    XH.VERSION_TAG = VERSION; XH.HOME_SHEET = S0; XH.HOME_LABEL = "Portada"
    XH.FOOTER_TEXT = f"Exergy EXG S.A.S. · Modelo FV Montecristi → GPM · {VERSION}"
    XH.META_LIVE = True
    wb = Workbook()
    wb.remove(wb.active)
    ws4 = build_energia(wb)
    ws5 = build_capex(wb)
    ws6 = build_opex(wb)
    ws7 = build_fiscal(wb)
    cases = build_cases()
    wsm = build_motor(wb, cases)
    ws10, tornado_rows = build_sensibilidad(wb)   # v3.0: antes de 08/09, que enlazan a 10 §A (V3)
    ws8 = build_flujo(wb)
    ws9 = build_exergy(wb)
    ws1 = build_supuestos(wb)   # tras 10: usa las anclas de sección para «Sensibilizado en»
    ws2 = build_legal(wb)
    ws3 = build_tramites(wb)
    ws11 = build_riesgos(wb)
    terms = "+".join([f"ABS('{SM}'!$B${brow('FCF_u', t)}-'{S8}'!{C(t)}{FLUJO['fcf']})" for t in TS])
    ws12 = build_fuentes(wb, terms)
    ws13 = build_controles(wb, terms)
    ws0 = build_portada(wb, ws8, ws7, ws5, ws10, tornado_rows)
    wsg = build_guia(wb)
    order = [S0, SG, S1, S2, S3, S4, S5, S6, S7, S8, S9, S10, S11, S12, S13, SM]
    wb._sheets = [wb[n] for n in order]
    wb.active = 0
    for n in order:
        wb[n].sheet_view.zoomScale = 90
    # v3.0 (V-b): ningún literal de texto > 255 caracteres (Excel los reescribiría como _xlfn._LONGTEXT)
    n_split = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and XH.long_literals(c.value):
                    c.value = XH.split_long_literals(c.value); n_split += 1
    print(f"  literales > 255 partidos: {n_split}")
    wb.calculation.fullCalcOnLoad = True
    from openpyxl.utils import get_column_letter as _gcl
    for ws in wb.worksheets:
        if not ws.print_area:
            ws.print_area = f"A1:{_gcl(ws.max_column)}{ws.max_row}"
    wb.properties.creator = "Exergy EXG S.A.S. · Claude"
    wb.properties.title = f"Modelo FV Montecristi → GPM {VERSION}"
    wb.save(OUT)
    print("saved", OUT)


if __name__ == "__main__":
    main()

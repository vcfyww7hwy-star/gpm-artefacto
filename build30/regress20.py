# -*- coding: utf-8 -*-
"""regress20.py — regresión numérica v2.0 frente a una referencia (ambos libros recalculados por LibreOffice).
Uso: regress20.py REF_calc.xlsx NEW_calc.xlsx [--tol 0] [--quiet] [--same]
 --same : ambos libros son v2.0 (p. ej. modelo completo vs Resumen unificado): casos con el mismo nombre, sin mapa de nombres;
          los nombres de REF ausentes en NEW se informan pero no son problema (el Resumen es un subconjunto del modelo).

 1. Nombres definidos: los de REF se traducen al nombre v2.0 (NAME_MAP; los de EXPECTED_GONE desaparecen por diseño) y su valor
    (celda o rango, celda a celda) se compara con tolerancia `tol`. Se listan los nombres nuevos y los ausentes no previstos.
 2. Motor: cada columna de REF se empareja por NOMBRE de caso (CASE_MAP para los renombrados; los casos sin homólogo se listan) y
    cada fila por número; las filas nuevas de v2.0 (24 kfix, 43 fKeff) se saltan. Las filas de parámetros (5–23) se informan aparte
    (pueden diferir por diseño: el Favorable v2.0 lleva fK = 1 y kfix = 0,75 en vez de fK = comercial/base); escalares, salidas y
    bloques anuales deben coincidir dentro de `tol`.
 3. Controles de NEW: N_Controles_OK = N_Controles y Estado_Controles empieza por «●».
Sale con 1 si hay alguna diferencia fuera de lo previsto."""
import sys, re
from openpyxl import load_workbook

ref_path, new_path = sys.argv[1], sys.argv[2]
tol = float(sys.argv[sys.argv.index("--tol") + 1]) if "--tol" in sys.argv else 0.0
quiet = "--quiet" in sys.argv
same = "--same" in sys.argv
MOTOR = "Motor_Sens"
# nombre v1.3 → nombre v2.0
NAME_MAP = {}
for key in ["E1", "Ahorro1", "TIR", "VAN", "PB", "LCOE", "TIReq", "VANeq", "DSCR", "Aporte", "VANX"]:
    for tag in "CBF":
        NAME_MAP[f"ESC_{key}_{tag}"] = f"{tag}_{key}"
EXPECTED_GONE = {"Escenario_Activo", "CAPEX_Caso", "CAPEX_Comercial_Wp", "CAPEX_Comercial", "Delta_Conservador", "Cons_OPEX_Up", "Cons_Peaje", "Fav_EscTarifa",
                 "Estado_Base", "N_Fuera_Base"}
# caso v1.3 → caso v2.0 (los demás conservan el nombre); None = desaparece por diseño
CASE_MAP = {"Base (activo)": "Custom", "P50": "Custom P50", "P90": "Custom P90", "Escenario Base": "Base", "Escenario Conservador": "Conservador",
            "Escenario Favorable": "Favorable", "P1": None, "P2": None, "P3": None, "P4": "P1", "P5": "P2", "P6": "P3", "P7": "P4"}
if same:
    NAME_MAP, EXPECTED_GONE, CASE_MAP = {}, set(), {}
NEW_ROWS = set() if same else {24, 43}
PARAM_ROWS = range(5, 24)

R = load_workbook(ref_path, data_only=True)
N = load_workbook(new_path, data_only=True)
RF = load_workbook(ref_path, data_only=False)
NF = load_workbook(new_path, data_only=False)
problems, info = [], []


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def norm(v):
    return int(v) if isinstance(v, bool) else v


def name_cells(wbv, wbf, nm):
    dn = wbf.defined_names.get(nm)
    if dn is None:
        return None
    ref = dn.attr_text.replace("'", "").replace("$", "")
    sh, rng = ref.split("!")
    ws = wbv[sh]
    if ":" in rng:
        out = []
        for row in ws[rng]:
            for c in (row if isinstance(row, tuple) else (row,)):
                out.append(norm(c.value))
        return out
    return [norm(ws[rng].value)]


# 1. nombres
ref_names = set(RF.defined_names.keys())
new_names = set(NF.defined_names.keys())
translated = {nm: NAME_MAP.get(nm, nm) for nm in ref_names}
missing = sorted(nm for nm, t in translated.items() if t not in new_names and nm not in EXPECTED_GONE)
gone = sorted(nm for nm in ref_names if nm in EXPECTED_GONE)
added = sorted(new_names - set(translated.values()))
if missing and not same:
    problems.append(f"nombres ausentes en NEW (no previstos): {missing}")
elif missing:
    info.append(f"nombres de REF que NEW no define (subconjunto): {len(missing)}")
info.append(f"nombres: REF {len(ref_names)} · NEW {len(new_names)} · desaparecen por diseño {len(gone)} · nuevos {len(added)}")
if not quiet:
    info.append(f"  nuevos: {added}")
n_vals = 0
n_diff = 0
diffs_by_name = []
for nm in sorted(ref_names):
    t = translated[nm]
    if t not in new_names:
        continue
    a = name_cells(R, RF, nm); b = name_cells(N, NF, t)
    if a is None or b is None:
        problems.append(f"nombre {nm}: no legible"); continue
    if len(a) != len(b):
        problems.append(f"nombre {nm}→{t}: tamaño distinto {len(a)} vs {len(b)}"); continue
    dmax = 0.0
    for va, vb in zip(a, b):
        n_vals += 1
        if is_num(va) and is_num(vb):
            d = abs(va - vb)
            if d > tol:
                n_diff += 1; dmax = max(dmax, d)
        elif is_num(va) != is_num(vb):
            n_diff += 1; dmax = float("inf")
        elif isinstance(va, str) and isinstance(vb, str) and va != vb:
            da, db = re.findall(r"[\d.,]+", va), re.findall(r"[\d.,]+", vb)
            if da != db and not quiet:
                info.append(f"  texto con cifras distinto: {nm}→{t}: {va[:70]!r} → {vb[:70]!r}")
    if dmax:
        diffs_by_name.append((nm, t, dmax))
info.append(f"nombres comparados: {len(ref_names) - len(gone) - len(missing)} ({n_vals} valores) · valores distintos {n_diff}")
for nm, t, dmax in diffs_by_name:
    problems.append(f"nombre {nm}→{t}: Δmax = {dmax:.6g}")

# 2. Motor por nombre de caso
wr, wn = R[MOTOR], N[MOTOR]
ref_cases = {}
for c in range(2, wr.max_column + 1):
    nm = wr.cell(row=4, column=c).value
    if nm:
        ref_cases[nm] = c
new_cases = {}
for c in range(2, wn.max_column + 1):
    nm = wn.cell(row=4, column=c).value
    if nm:
        new_cases[nm] = c
unmatched, dropped = [], []
n_cmp = n_bad = n_param_diff = 0
worst = []
param_notes = []
for nm, c in ref_cases.items():
    t = CASE_MAP.get(nm, nm)
    if t is None:
        dropped.append(nm); continue
    if t not in new_cases:
        unmatched.append(f"{nm}→{t}"); continue
    c2 = new_cases[t]
    for r in range(5, wr.max_row + 1):
        if r in NEW_ROWS:
            continue
        va = wr.cell(row=r, column=c).value
        if not is_num(va):
            continue
        vb = wn.cell(row=r, column=c2).value
        if r in PARAM_ROWS:
            if not is_num(vb) or abs(va - vb) > tol:
                n_param_diff += 1
                param_notes.append(f"{nm}→{t} fila {r} ({wr.cell(row=r, column=1).value}): {va!r} → {vb!r}")
            continue
        n_cmp += 1
        if not is_num(vb) or abs(va - vb) > tol:
            n_bad += 1
            worst.append((abs(va - vb) if is_num(vb) else float("inf"), f"{nm}→{t} fila {r} ({wr.cell(row=r, column=1).value}): {va!r} → {vb!r}"))
info.append(f"Motor: casos REF {len(ref_cases)} · NEW {len(new_cases)} · emparejados {len(ref_cases) - len(dropped) - len(unmatched)} · desaparecen por diseño {dropped}")
info.append(f"Motor: {n_cmp} celdas (escalares, salidas, bloques) comparadas · distintas {n_bad} · parámetros distintos (informativo) {n_param_diff}")
if not quiet:
    for p in param_notes[:12]:
        info.append("  param: " + p)
if unmatched:
    problems.append(f"Motor: casos sin homólogo: {unmatched}")
if n_bad:
    worst.sort(reverse=True)
    problems.append(f"Motor: {n_bad} celdas distintas; peores: " + " | ".join(w for _, w in worst[:8]))

# 3. controles
try:
    ok = name_cells(N, NF, "N_Controles_OK")[0]; tot = name_cells(N, NF, "N_Controles")[0]; est = name_cells(N, NF, "Estado_Controles")[0]
    info.append(f"controles NEW: {ok}/{tot} · {est}")
    if ok != tot or not str(est).startswith("●"):
        problems.append(f"controles: {ok}/{tot} · {est}")
except Exception as e:
    info.append(f"controles: no evaluables ({e})")

print(f"regress20 · REF={ref_path} · NEW={new_path} · tol={tol}")
for x in info:
    print("  ·", x)
if problems:
    print(f"  ✖ {len(problems)} problema(s):")
    for p in problems:
        print("    -", p)
    sys.exit(1)
print("  ✔ regresión sin diferencias fuera de lo previsto")

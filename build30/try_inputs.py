# -*- coding: utf-8 -*-
"""try_inputs.py RAW.xlsx SALIDA_calc.xlsx Nombre=valor|Hoja!Celda=valor [… ] [--shadow] [--quiet]
Prueba de uso con LibreOffice: escribe valores en celdas con nombre del libro (01_Supuestos; ""/vacio = borrar), recalcula, e imprime
los KPI de los cuatro casos, los controles que no están en ● y (opcional) la sombra. Sirve para validar los M en modo activo y para la
prueba de uso F4 antes de repetirla en Excel real. Valores: números (0.98, 200000), textos (SALELGI), vacio → None."""
import sys, subprocess, os, shutil
from openpyxl import load_workbook

raw, out = sys.argv[1], sys.argv[2]
sets = [a for a in sys.argv[3:] if "=" in a]
quiet = "--quiet" in sys.argv; do_shadow = "--shadow" in sys.argv
wb = load_workbook(raw)


def dest(nm):
    if "!" in nm:   # Hoja!Celda (p. ej. 03_Tramites!G21) — celdas de diseño sin nombre
        sh, ref = nm.split("!", 1)
        return wb[sh.strip("'")], ref.replace("$", "")
    dn = wb.defined_names[nm]
    for sh, ref in dn.destinations:
        return wb[sh], ref.replace("$", "")


def parse(v):
    if v.lower() in ("vacio", "vacío", "none", '""', ""):
        return None
    try:
        return int(v) if v.lstrip("-").isdigit() else float(v)
    except ValueError:
        return v


changes = []
for a in sets:
    nm, v = a.split("=", 1)
    ws, ref = dest(nm)
    old = ws[ref].value
    ws[ref].value = parse(v)
    changes.append((nm, old, parse(v)))
wb.save(out)
subprocess.run(["python3", "/mnt/skills/public/xlsx/scripts/recalc.py", out, "300"], capture_output=True, text=True)
wv = load_workbook(out, data_only=True)


def val(n):
    dn = wv.defined_names.get(n)
    if dn is None:
        return None
    for sh, ref in dn.destinations:
        return wv[sh][ref.replace("$", "")].value


def f(x, fmt):
    if isinstance(x, (int, float)):
        return format(x, fmt)
    return str(x)


print("cambios:", "; ".join(f"{nm}: {o!r} → {n!r}" for nm, o, n in changes))
print(f"  {val('Estado_Controles')} | {val('Estado_Custom')} | {val('Estado_Neutro')} | {val('Estado_Entregado')} | Meses_Construccion={val('Meses_Construccion')} · Potencia_AC={val('Potencia_AC')} · N_Por_Confirmar={val('N_Por_Confirmar')}")
for tag in "XCBF":
    print(f"  {tag}: TIR {f(val(tag+'_TIR'), '.4%')} · VAN {f(val(tag+'_VAN'), ',.0f')} · acc {f(val(tag+'_TIReq'), '.4%')} · VANeq {f(val(tag+'_VANeq'), ',.0f')} · DSCR {f(val(tag+'_DSCR'), '.4f')} · PB {f(val(tag+'_PB'), '.2f')} · LCOE {f(val(tag+'_LCOE'), '.2f')} · VANX {f(val(tag+'_VANX'), ',.0f')}")
ws13 = wv["13_Controles"]
for row in ws13.iter_rows(min_row=5):
    if row[2].value and row[3].value and not str(row[3].value).startswith("●"):
        print("   ", row[2].value, str(row[3].value)[:110])
if not quiet:
    for row in ws13.iter_rows(min_row=5):
        if row[2].value and str(row[2].value).startswith(("H", "I")) and row[3].value:
            print("   ", row[2].value, str(row[3].value)[:110])
if do_shadow:
    r = subprocess.run(["python3", os.path.join(os.path.dirname(os.path.abspath(__file__)), "shadow30.py"), out, out.replace(".xlsx", "_sombra.csv"), "--blocks", "--quiet"], capture_output=True, text=True)
    print("  " + r.stdout.strip().replace("\n", "\n  "))

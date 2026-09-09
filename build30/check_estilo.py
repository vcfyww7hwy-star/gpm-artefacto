# -*- coding: utf-8 -*-
"""check_estilo.py — cumplimiento del sistema visual v1.3 sobre un .xlsx.
Falla (exit 1) si aparece: un color fuera de la paleta (fuente, relleno, borde, formato condicional, pestaña, gráfico),
un panel congelado, una fuente distinta de Calibri / Calibri Light, un tamaño < 8,5 pt (03_Tramites admite ≥ 8 en el Gantt),
o un gráfico sin la plantilla v2 (style=2, sin varyColors, ejes ≥ 8,5 pt, colores de paleta).
Uso: check_estilo.py libro.xlsx [--allow-size-sheet HOJA ...] [--max 25]"""
import sys, re, zipfile
from collections import Counter, defaultdict
from openpyxl import load_workbook
from xl_helpers import PALETTE, PALETTE_NAMES, FONTS_OK, MIN_PT, CHART_STYLE_ID, long_literals

path = sys.argv[1]
max_show = 25
allow_size_sheets = {"03_Tramites"}
for i, a in enumerate(sys.argv):
    if a == "--max":
        max_show = int(sys.argv[i + 1])
    if a == "--allow-size-sheet":
        allow_size_sheets.add(sys.argv[i + 1])

findings = defaultdict(list)   # tipo → [(hoja, celda, detalle)]


def rgb6(c):
    """Devuelve RRGGBB o None si no es un color explícito (theme/indexed/None)."""
    if c is None:
        return None
    rgb = getattr(c, "rgb", None)
    if not isinstance(rgb, str):
        return None
    rgb = rgb.upper()
    if len(rgb) == 8:
        rgb = rgb[2:]
    return rgb


def chk_color(kind, sheet, cell, rgb):
    if rgb is None or rgb == "000000" and kind == "borde":
        return
    if rgb not in PALETTE:
        findings[f"color fuera de paleta ({kind})"].append((sheet, cell, rgb))


wb = load_workbook(path)
n_cells = 0
for ws in wb.worksheets:
    t = ws.title
    if ws.freeze_panes:
        findings["panel congelado"].append((t, ws.freeze_panes, ""))
    if ws.sheet_view.pane is None and any(sel.pane for sel in ws.sheet_view.selection):
        findings["selección con pane sin <pane> (Excel pide reparar)"].append((t, [sel.pane for sel in ws.sheet_view.selection], ""))
    tab = ws.sheet_properties.tabColor
    if tab is not None:
        rgb = rgb6(tab)
        if rgb is not None and rgb not in PALETTE:
            findings["color fuera de paleta (pestaña)"].append((t, "tab", rgb))
    min_pt = 8.0 if t in allow_size_sheets else MIN_PT
    for row in ws.iter_rows():
        for c in row:
            if c.value is None and not c.has_style:
                continue
            n_cells += 1
            if isinstance(c.value, str) and c.value.startswith("=") and long_literals(c.value):
                findings["literal de texto > 255 caracteres (Excel → _xlfn._LONGTEXT)"].append((t, c.coordinate, f"{max(len(x) for x in long_literals(c.value))} car."))
            f = c.font
            if f is not None:
                if f.name and f.name not in FONTS_OK:
                    findings["fuente distinta de Calibri"].append((t, c.coordinate, f.name))
                if f.sz is not None and c.value is not None and float(f.sz) < min_pt:
                    findings["tamaño < mínimo"].append((t, c.coordinate, f"{f.sz} pt"))
                chk_color("fuente", t, c.coordinate, rgb6(f.color))
            fl = c.fill
            if fl is not None and fl.fill_type == "solid":
                chk_color("relleno", t, c.coordinate, rgb6(fl.fgColor))
            b = c.border
            if b is not None:
                for side_name in ("left", "right", "top", "bottom"):
                    sd = getattr(b, side_name)
                    if sd is not None and sd.style:
                        chk_color("borde", t, c.coordinate, rgb6(sd.color))
    # formato condicional
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            if rule.type == "colorScale":
                findings["escala de color (mapa de calor con relleno)"].append((t, str(cf.sqref), rule.type))
            d = rule.dxf
            if d is None:
                continue
            if d.font is not None:
                chk_color("cf-fuente", t, str(cf.sqref), rgb6(d.font.color))
            if d.fill is not None:
                fg = d.fill.fgColor if hasattr(d.fill, "fgColor") else None
                bg = d.fill.bgColor if hasattr(d.fill, "bgColor") else None
                for cc in (fg, bg):
                    r = rgb6(cc)
                    if r is not None and r not in PALETTE:
                        findings["color fuera de paleta (cf-relleno)"].append((t, str(cf.sqref), r))
            if d.border is not None:
                for side_name in ("left", "right", "top", "bottom"):
                    sd = getattr(d.border, side_name)
                    if sd is not None and sd.style:
                        chk_color("cf-borde", t, str(cf.sqref), rgb6(sd.color))

# ---- gráficos (XML)
z = zipfile.ZipFile(path)
charts = sorted(n for n in z.namelist() if re.match(r"xl/charts/chart\d+\.xml$", n))
# hoja de cada gráfico: drawing → sheet
rels_sheet = {}
try:
    wbxml = z.read("xl/workbook.xml").decode("utf8")
    wbrels = z.read("xl/_rels/workbook.xml.rels").decode("utf8")
    rid_target = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="([^"]+)"', wbrels))
    rid_target.update({b: a for a, b in re.findall(r'Target="([^"]+)"[^>]*Id="(rId\d+)"', wbrels)})
    for nm, rid in re.findall(r'<sheet [^>]*name="([^"]+)"[^>]*r:id="(rId\d+)"', wbxml):
        tgt = rid_target.get(rid, "")
        sheet_file = "xl/" + tgt if not tgt.startswith("/") else tgt[1:]
        rel_file = sheet_file.replace("worksheets/", "worksheets/_rels/") + ".rels"
        if rel_file in z.namelist():
            for d in re.findall(r'Target="(?:/xl/|\.\./)drawings/(drawing\d+\.xml)"', z.read(rel_file).decode("utf8")):
                drel = f"xl/drawings/_rels/{d}.rels"
                if drel in z.namelist():
                    for chn in re.findall(r'Target="(?:/xl/|\.\./)charts/(chart\d+\.xml)"', z.read(drel).decode("utf8")):
                        rels_sheet[f"xl/charts/{chn}"] = nm.replace("&amp;", "&")
except Exception as e:
    pass
for chn in charts:
    xml = z.read(chn).decode("utf8")
    sheet = rels_sheet.get(chn, "?")
    tag = f"{sheet} · {chn.split('/')[-1]}"
    cols = set(re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"', xml))
    bad = {c.upper() for c in cols} - PALETTE
    if bad:
        findings["color fuera de paleta (gráfico)"].append((sheet, chn.split('/')[-1], ",".join(sorted(bad))))
    if not re.search(rf'<(?:c:)?style val="{CHART_STYLE_ID}"/>', xml):
        findings["gráfico sin plantilla v2 (style)"].append((sheet, chn.split('/')[-1], ""))
    if re.search(r'<(?:c:)?varyColors val="1"/>', xml):
        findings["gráfico sin plantilla v2 (varyColors)"].append((sheet, chn.split('/')[-1], ""))
    for sz in re.findall(r'sz="(\d+)"', xml):
        if int(sz) < int(MIN_PT * 100):
            findings["gráfico: texto < 8,5 pt"].append((sheet, chn.split('/')[-1], f"{int(sz)/100} pt"))
            break
    if not re.search(r'<(?:c:)?title>', xml):
        findings["gráfico sin título"].append((sheet, chn.split('/')[-1], ""))

# ---- informe
total = sum(len(v) for v in findings.values())
print(f"check_estilo · {path}")
print(f"  hojas: {len(wb.worksheets)} · celdas con estilo: {n_cells} · gráficos: {len(charts)}")
if total == 0:
    print("  ✔ 0 hallazgos")
    sys.exit(0)
for k, v in sorted(findings.items(), key=lambda kv: -len(kv[1])):
    print(f"  ✖ {k}: {len(v)}")
    by_sheet = Counter(x[0] for x in v)
    print("     por hoja: " + ", ".join(f"{s}={n}" for s, n in by_sheet.most_common()))
    for x in v[:max_show]:
        print(f"     {x[0]}!{x[1]} {x[2]}")
print(f"  TOTAL hallazgos: {total}")
sys.exit(1)

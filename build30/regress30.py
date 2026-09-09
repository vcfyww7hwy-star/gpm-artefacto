# -*- coding: utf-8 -*-
"""regress30.py — regresión numérica v3.0 frente a una referencia (ambos libros recalculados por LibreOffice).
Uso: regress30.py REF_calc.xlsx NEW_calc.xlsx [--tol 0] [--quiet] [--same] [--allow n1,n2,…] [--gone n1,n2,…] [--rows-gone e1;e2;…]
 --same       : NEW es un subconjunto de REF (p. ej. modelo completo vs Resumen unificado): los nombres de REF ausentes en NEW se informan
                pero no son problema; las filas del Motor ausentes en NEW tampoco.
 --allow      : nombres definidos que pueden diferir por diseño (estructurales: N_Controles, N_Por_Confirmar_Esperado, …); se informan con sus valores.
 --gone       : nombres de REF que desaparecen por diseño en NEW.
 --rows-gone  : etiquetas de fila del Motor (columna A) de REF que desaparecen por diseño en NEW.

 1. Nombres definidos: cada nombre de REF presente en NEW se compara valor a valor (celda o rango) con tolerancia `tol`; se listan los nuevos
    (informativo) y los ausentes no previstos (problema, salvo --same/--gone).
 2. Motor (Motor_Sens): cada columna de REF se empareja por NOMBRE de caso (fila 4) y cada fila por ETIQUETA (columna A: texto del parámetro/
    escalar/salida; «Bloque X» + t para las filas anuales), no por número de fila → las inserciones del Motor v3.0 no rompen la comparación.
    Las filas de parámetros se informan aparte (pueden diferir por diseño); escalares, salidas y bloques deben coincidir dentro de `tol`.
    Filas y casos nuevos de NEW: informativos. Filas de REF ausentes en NEW: problema (salvo --same/--rows-gone).
 3. Controles de NEW: N_Controles_OK = N_Controles y Estado_Controles empieza por «●».
Sale con 1 si hay alguna diferencia fuera de lo previsto."""
import sys, re
from openpyxl import load_workbook

ref_path, new_path = sys.argv[1], sys.argv[2]


def arg(flag, default=None):
    return sys.argv[sys.argv.index(flag) + 1] if flag in sys.argv else default


tol = float(arg("--tol", "0"))
quiet = "--quiet" in sys.argv
same = "--same" in sys.argv
ALLOW = set(filter(None, arg("--allow", "").split(",")))
GONE = set(filter(None, arg("--gone", "").split(",")))
ROWS_GONE = set(filter(None, arg("--rows-gone", "").split(";")))
MOTOR = "Motor_Sens"
# etiquetas de las filas de parámetros (v2.0 + v3.0): se informan, no se exigen iguales (los casos pueden definirse distinto por diseño)
PARAM_LABELS = {"Factor CAPEX (× bottom-up)", "CAPEX fijo $/Wp (0 = bottom-up × factor)", "Factor tarifa", "Escenario energía (1=P50, 2=P90)", "Factor OPEX", "Peaje $/kWh",
                "Escalación tarifa", "Participación (1/0)", "IVA recuperable (1/0)", "Contrato inversión (1/0)", "Deuda (1/0)", "Tasa deuda", "Apalancamiento", "Plazo", "Gracia",
                "Potencia DC (kWp)", "Ratio DC/AC", "Terreno lo compra SALELGI (1/0)", "Banco financia terreno (1/0)", "Factor precio terreno",
                "Disponibilidad", "Escalación CAPEX (%/año hasta la compra)", "Peaje por potencia $/kW-mes", "Reemplazo inversores (0 No · 1 SALELGI · 2 Exergy)",
                "Utilidad gravable SALELGI (−1 = ilimitada)", "Meses de construcción", "Tasa descuento accionista", "Desmantelamiento (% CAPEX en t=H)", "Degradación adicional (%/año)",
                "Deducción adicional aplicable (1/0)"}   # v3.1

R = load_workbook(ref_path, data_only=True)
N = load_workbook(new_path, data_only=True)
RF = load_workbook(ref_path, data_only=False)
NF = load_workbook(new_path, data_only=False)
problems, info = [], []


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def norm(v):
    return int(v) if isinstance(v, bool) else v


def name_cells(wbv, wbf, nm):
    dn = wbf.defined_names.get(nm)
    if dn is None:
        return None
    ref = dn.attr_text.replace("'", "").replace("$", "")
    sh, rng = ref.split("!")
    ws = wbv[sh]
    if ":" in rng:
        out = []
        for row in ws[rng]:
            for c in (row if isinstance(row, tuple) else (row,)):
                out.append(norm(c.value))
        return out
    return [norm(ws[rng].value)]


# 1. nombres
ref_names = set(RF.defined_names.keys())
new_names = set(NF.defined_names.keys())
missing = sorted(nm for nm in ref_names if nm not in new_names and nm not in GONE)
gone = sorted(nm for nm in ref_names if nm in GONE)
added = sorted(new_names - ref_names)
if missing and not same:
    problems.append(f"nombres ausentes en NEW (no previstos): {missing}")
elif missing:
    info.append(f"nombres de REF que NEW no define (subconjunto): {len(missing)}")
info.append(f"nombres: REF {len(ref_names)} · NEW {len(new_names)} · desaparecen por diseño {len(gone)} · nuevos {len(added)}")
if not quiet and added:
    info.append(f"  nuevos: {added}")
n_vals = n_diff = 0
diffs_by_name = []; allowed_diffs = []
for nm in sorted(ref_names & new_names):
    a = name_cells(R, RF, nm); b = name_cells(N, NF, nm)
    if a is None or b is None:
        problems.append(f"nombre {nm}: no legible"); continue
    if len(a) != len(b):
        if nm in ALLOW:
            allowed_diffs.append(f"{nm}: tamaño {len(a)} → {len(b)}")
        else:
            problems.append(f"nombre {nm}: tamaño distinto {len(a)} vs {len(b)}")
        continue
    dmax = 0.0
    for va, vb in zip(a, b):
        n_vals += 1
        if is_num(va) and is_num(vb):
            d = abs(va - vb)
            if d > tol:
                dmax = max(dmax, d)
                if nm not in ALLOW: n_diff += 1
        elif is_num(va) != is_num(vb):
            dmax = float("inf")
            if nm not in ALLOW: n_diff += 1
        elif isinstance(va, str) and isinstance(vb, str) and va != vb:
            da, db = re.findall(r"[\d.,]+", va), re.findall(r"[\d.,]+", vb)
            if da != db and not quiet:
                info.append(f"  texto con cifras distinto: {nm}: {va[:70]!r} → {vb[:70]!r}")
    if dmax:
        if nm in ALLOW:
            allowed_diffs.append(f"{nm}: {a if len(a) <= 4 else str(len(a)) + ' valores'} → {b if len(b) <= 4 else str(len(b)) + ' valores'}")
        else:
            diffs_by_name.append((nm, dmax))
info.append(f"nombres comparados: {len(ref_names & new_names)} ({n_vals} valores) · valores distintos {n_diff}" + (f" · permitidos por diseño {len(allowed_diffs)}" if allowed_diffs else ""))
for x in allowed_diffs:
    info.append("  permitido: " + x)
for nm, dmax in diffs_by_name:
    problems.append(f"nombre {nm}: Δmax = {dmax:.6g}")


# 2. Motor por nombre de caso y etiqueta de fila
def motor_rows(ws):
    rows = {}; block = None
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v is None:
            continue
        if isinstance(v, str) and v.startswith("Bloque "):
            block = v.split(" ")[1]; continue
        if block is not None and is_num(v):
            rows[(block, int(v))] = r
        elif isinstance(v, str) and r >= 5:
            rows[v] = r
    return rows


def case_cols(ws):
    out = {}
    for c in range(2, ws.max_column + 1):
        nm = ws.cell(row=4, column=c).value
        if nm:
            out[nm] = c
    return out


wr, wn = R[MOTOR], N[MOTOR]
rr, rn = motor_rows(wr), motor_rows(wn)
ref_cases, new_cases = case_cols(wr), case_cols(wn)
rows_missing = sorted((k for k in rr if k not in rn and (k if isinstance(k, str) else k[0]) not in ROWS_GONE), key=str)
rows_new = sorted((k for k in rn if k not in rr), key=str)
rows_gone = sorted((k for k in rr if k not in rn and (k if isinstance(k, str) else k[0]) in ROWS_GONE), key=str)
common_rows = [k for k in rr if k in rn]
if rows_missing and not same:
    problems.append(f"Motor: {len(rows_missing)} fila(s) de REF ausentes en NEW: {rows_missing[:12]}")
elif rows_missing:
    info.append(f"Motor: filas de REF que NEW no tiene (subconjunto): {len(rows_missing)}")
new_scalar_rows = [k for k in rows_new if isinstance(k, str)]
new_blocks = sorted({k[0] for k in rows_new if isinstance(k, tuple)})
info.append(f"Motor: filas REF {len(rr)} · NEW {len(rn)} · comunes {len(common_rows)} · nuevas en NEW {len(rows_new)}" + (f" ({new_scalar_rows} + bloques {new_blocks})" if rows_new else "") + (f" · desaparecen por diseño {len(rows_gone)}" if rows_gone else ""))
unmatched = [nm for nm in ref_cases if nm not in new_cases]
cases_new = [nm for nm in new_cases if nm not in ref_cases]
n_cmp = n_bad = n_param_diff = 0
worst = []; param_notes = []
for nm, c in ref_cases.items():
    if nm not in new_cases:
        continue
    c2 = new_cases[nm]
    for k in common_rows:
        va = wr.cell(row=rr[k], column=c).value
        if not is_num(va):
            continue
        vb = wn.cell(row=rn[k], column=c2).value
        if isinstance(k, str) and k in PARAM_LABELS:
            if not is_num(vb) or abs(va - vb) > tol:
                n_param_diff += 1
                param_notes.append(f"{nm} · {k}: {va!r} → {vb!r}")
            continue
        n_cmp += 1
        if not is_num(vb) or abs(va - vb) > tol:
            n_bad += 1
            worst.append((abs(va - vb) if is_num(vb) else float("inf"), f"{nm} · {k}: {va!r} → {vb!r}"))
info.append(f"Motor: casos REF {len(ref_cases)} · NEW {len(new_cases)} · emparejados {len(ref_cases) - len(unmatched)} · nuevos en NEW {len(cases_new)}" + (f" {cases_new}" if cases_new and not quiet else ""))
info.append(f"Motor: {n_cmp} celdas (escalares, salidas, bloques) comparadas por etiqueta · distintas {n_bad} · parámetros distintos (informativo) {n_param_diff}")
if not quiet:
    for p in param_notes[:12]:
        info.append("  param: " + p)
if unmatched:
    problems.append(f"Motor: casos de REF sin homólogo en NEW: {unmatched}")
if n_bad:
    worst.sort(reverse=True)
    problems.append(f"Motor: {n_bad} celdas distintas; peores: " + " | ".join(w for _, w in worst[:8]))

# 3. controles
try:
    ok = name_cells(N, NF, "N_Controles_OK")[0]; tot = name_cells(N, NF, "N_Controles")[0]; est = name_cells(N, NF, "Estado_Controles")[0]
    info.append(f"controles NEW: {ok}/{tot} · {est}")
    if ok != tot or not str(est).startswith("●"):
        problems.append(f"controles: {ok}/{tot} · {est}")
except Exception as e:
    info.append(f"controles: no evaluables ({e})")

print(f"regress30 · REF={ref_path} · NEW={new_path} · tol={tol}" + (" · --same" if same else ""))
for x in info:
    print("  ·", x)
if problems:
    print(f"  ✖ {len(problems)} problema(s):")
    for p in problems:
        print("    -", p)
    sys.exit(1)
print("  ✔ regresión sin diferencias fuera de lo previsto")

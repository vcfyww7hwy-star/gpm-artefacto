# -*- coding: utf-8 -*-
"""check_errors.py LIBRO_calc.xlsx [--allow "Hoja!A1:B2,Hoja!C5"]
Lista las celdas con error (#N/A, #DIV/0!, #REF!, #VALUE!, #NAME?, #NUM!, #NULL!) en un libro recalculado
(valores en caché) y falla si alguna cae fuera de los rangos permitidos.
Los únicos errores admitidos en el modelo son los NA() intencionales de las series auxiliares de la portada
(marcadores del año de cruce: Excel no dibuja los puntos #N/A)."""
import sys
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries, get_column_letter

ERR = ("#N/A", "#DIV/0!", "#REF!", "#VALUE!", "#NAME?", "#NUM!", "#NULL!")
path = sys.argv[1]
allow = []
if "--allow" in sys.argv:
    for spec in sys.argv[sys.argv.index("--allow") + 1].split(","):
        sh, rng = spec.split("!")
        allow.append((sh, range_boundaries(rng)))


def allowed(sh, r, c):
    for s, (c1, r1, c2, r2) in allow:
        if s == sh and r1 <= r <= r2 and c1 <= c <= c2:
            return True
    return False


wb = load_workbook(path, data_only=True)
n_err = n_ok = 0
bad = []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith("#") and v in ERR:
                n_err += 1
                if allowed(ws.title, cell.row, cell.column):
                    n_ok += 1
                else:
                    bad.append(f"{ws.title}!{cell.coordinate} = {v}")
print(f"celdas con error: {n_err} · admitidas (NA() intencionales): {n_ok} · fuera de lista: {len(bad)}")
for b in bad[:20]:
    print("   ✖", b)
if bad:
    sys.exit(1)
print("✔ sin errores fuera de los rangos permitidos")

"""compara_raiz.py ENTREGADO.xlsx RAIZ.xlsx — compara (sólo lectura) entradas con nombre, nombres desplazados y fórmulas por hoja."""
import sys
from openpyxl import load_workbook
a,b=sys.argv[1],sys.argv[2]
A=load_workbook(a,data_only=True);B=load_workbook(b,data_only=True);Af=load_workbook(a);Bf=load_workbook(b)
def one(wbf,name):
    dn=wbf.defined_names.get(name)
    if dn is None: return None
    d=list(dn.destinations)
    if len(d)!=1: return None
    sh,ref=d[0]; ref=ref.replace("$","")
    return None if ":" in ref else (sh,ref)
isf=lambda x: isinstance(x,str) and x.startswith("=")
moved=[];diffs=[]
for n in Af.defined_names:
    ra=one(Af,n); rb=one(Bf,n)
    if not ra or not rb: continue
    if ra!=rb: moved.append((n,ra,rb))
    fa,fb=Af[ra[0]][ra[1]].value,Bf[rb[0]][rb[1]].value
    va,vb=A[ra[0]][ra[1]].value,B[rb[0]][rb[1]].value
    if not isf(fa) and not isf(fb):
        if va!=vb: diffs.append((n,ra,va,vb))
    elif isf(fa)!=isf(fb): diffs.append((n,ra,fa if isf(fa) else va,fb if isf(fb) else vb))
print("nombres desplazados:",len(moved)); print("entradas distintas:",len(diffs))
for d in diffs: print("  ",d)
for sh in Af.sheetnames:
    fa=sum(1 for r in Af[sh].iter_rows() for c in r if isf(c.value)); fb=sum(1 for r in Bf[sh].iter_rows() for c in r if isf(c.value))
    print(f"  {sh:22s} fórmulas {fa:6d}/{fb:6d} filas {Af[sh].max_row:4d}/{Bf[sh].max_row:4d}")
